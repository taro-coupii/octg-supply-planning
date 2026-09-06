"""MRP -> .xlsx export. PRESENTATION ONLY.

The workbook this builds is the modern replacement for the source workbook's
"Material Order Req" (the summary of what to order) + "By Item" (the
justification behind each line) sheets, exported back OUT to Excel so a planner
can hand it to a customer.

NOT ONE NUMBER IS COMPUTED HERE
-------------------------------
Every figure comes from `app.engines.mrp.mrp_summary` and
`app.engines.mrp.by_item`, called exactly as the `/mrp/summary` and
`/mrp/by-item` routes call them. There is no lead-time arithmetic, no runout
month-walk and no recoverability test in this module, and there must never be:
this platform's whole claim is that the order date on the screen and the order
date in the customer's spreadsheet are the same number because they came from
the same call. A second implementation shaped to fit a cell is the exact way
that stops being true. `today` is threaded through to both engines rather than
re-read per call so the summary sheet and the supporting sheets describe one
moment.

SHEET LAYOUT: SUMMARY + ONE DEMAND-DETAIL SHEET PER PRODUCT
------------------------------------------------------------
Tab 1 is the summary. Then ONE SHEET PER RECOMMENDED PRODUCT ("D01", "D02",
...), each laid out like the By Item screen: the monthly ledger (opening /
demand / contingency split / ownership-split closing balances) followed by
the demand lines charged to that product. The remaining evidence kinds
(inventory / runout / lead time) stay as fixed cross-product sheets with a
Product column.

Per-product sheets were ONCE rejected here because sanitised-and-truncated
product descriptions collide (variants differ near the END of the string,
which truncation removes). The scheme below dodges that entirely: every name
begins with a unique ordinal ("D01 ", "D02 ", ...) so uniqueness is by
construction, and the description that follows is best-effort context only.
The 31-character cap and illegal-character rules are enforced by
`_item_sheet_name`, and uniqueness is asserted at build time.

HONESTY RULES ARE CARRIED INTO THE CELLS
----------------------------------------
`InventoryPosition.on_order` and `.customer_owned` are `float | None`, where
None means UNKNOWN (no row anywhere) and 0.0 means measured-and-nothing. The
screens render those two differently and so does this workbook: an unknown
figure becomes the text `NO_DATA`, never a numeric 0. A bare 0 in a
customer-facing spreadsheet is a claim, and it would be the wrong one.

Likewise every quantity column names its unit -- in the header where a sheet's
quantities are all one product's, and in an adjacent `Unit` column where they
are not (a substituted demand line is listed on the substitute's rows but is
quantified in its OWN product's unit; see `app.engines.mrp.ByItemDemandLine`).
An unlabelled number in a workbook that mixes MTR and PCS is the spreadsheet
equivalent of the bare scalar `tests/test_units_of_measure.py` refuses.

A lead time that is NOT MODELLED is never printed as "0 months" -- 0 reads as
instant delivery, which would make the least-known product look like the
safest row. It is printed as `NOT_MODELLED` with the missing dimensions named,
mirroring the frontend's `LeadTimeCell`.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime

from sqlalchemy.orm import Session

from app.engines.mrp import ByItemAnalysis, MrpRecommendation, by_item, mrp_summary

# ---------------------------------------------------------------------------
# Sheet names -- constants, so they cannot collide and cannot be truncated.
# ---------------------------------------------------------------------------

SHEET_SUMMARY = "MRP Summary"
SHEET_INVENTORY = "Inventory Position"
SHEET_RUNOUT = "Runout Projection"
SHEET_LEAD_TIME = "Lead Time Breakdown"
SHEET_NOTES = "Notes"

SHEET_NAMES = (
    SHEET_SUMMARY,
    SHEET_INVENTORY,
    SHEET_RUNOUT,
    SHEET_LEAD_TIME,
    SHEET_NOTES,
)

#: Excel's own restrictions, asserted below against the constants above.
EXCEL_SHEET_NAME_MAX = 31
EXCEL_SHEET_NAME_ILLEGAL = set(":\\/?*[]")

for _name in SHEET_NAMES:
    assert 0 < len(_name) <= EXCEL_SHEET_NAME_MAX, _name
    assert not (set(_name) & EXCEL_SHEET_NAME_ILLEGAL), _name
assert len(set(SHEET_NAMES)) == len(SHEET_NAMES), SHEET_NAMES

#: What an UNKNOWN figure looks like in a cell. Text, deliberately, so it can
#: never be mistaken for -- or summed with -- a measured 0.
NO_DATA = "no data (unknown)"
#: What a lead time the model cannot total looks like. Never "0".
NOT_MODELLED = "NOT MODELLED"

_ORDERABLE = "Order"
_ESCALATE = "ESCALATE - unrecoverable"


@dataclass
class MrpExport:
    content: bytes
    filename: str
    #: Recommendation rows written to tab 1, split the same way the sheet splits
    #: them, so a caller (and the round-trip test) can assert the file is not
    #: empty by accident without reopening it.
    recommendation_count: int
    orderable_count: int
    unrecoverable_count: int
    customer_name: str | None
    sheet_names: tuple[str, ...]
    #: Products whose supporting detail could NOT be built, with the reason --
    #: see `_analyses`. Reported rather than swallowed.
    unavailable_products: tuple[tuple[str, str], ...] = ()


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------


def _unit(value) -> str:
    """A UnitOfMeasure's display value, tolerant of a raw string."""
    return getattr(value, "value", value) or ""


def _qty(value: float | None) -> float | str:
    """A quantity cell. None is UNKNOWN and becomes text, NEVER 0.0.

    This is the whole reason the function exists: `openpyxl` would happily write
    `None` as an empty cell, and an empty numeric cell in a column of quantities
    is read by a human as zero just as reliably as a literal 0 is. The text says
    what is actually true.
    """
    return NO_DATA if value is None else float(value)


def _day(value: date | datetime | None) -> date | str:
    if value is None:
        return NO_DATA
    return value.date() if isinstance(value, datetime) else value


def _lead_time_cell(rec_or_analysis) -> tuple[float | str, str]:
    """(months cell, explanation) for a row carrying a `lead_time` breakdown.

    Returns `NOT_MODELLED` rather than the scalar whenever the breakdown says
    the model is incomplete, and names the missing dimensions. `lead_time` is
    None only on rows built before the breakdown existed; the scalar is then the
    only thing available and is reported with that caveat.
    """
    lead = getattr(rec_or_analysis, "lead_time", None)
    months = getattr(rec_or_analysis, "lead_time_months", None)
    if lead is None:
        if months:
            return float(months), "breakdown unavailable for this row"
        return NOT_MODELLED, "no lead-time breakdown available"
    if not lead.modelled:
        missing = ", ".join(lead.missing_dimensions) or "any dimension"
        return NOT_MODELLED, (
            f"missing lead-time components for: {missing}. Matched so far "
            f"{lead.matched_months:g} month(s) -- NOT a lead time, do not use it "
            "as one. Any order date shown is indicative only."
        )
    return float(lead.total_months), (
        " + ".join(
            f"{c.dimension} {c.attribute_value}{' (shared)' if c.shared else ''} "
            f"{c.months:g}mo"
            for c in lead.components
        )
        or "modelled"
    )


def _product_label(product_id: str, description: str | None) -> str:
    return description or product_id


# ---------------------------------------------------------------------------
# Styling -- minimal, matching the demand template's restraint
# ---------------------------------------------------------------------------


def _write_header(sheet, columns: list[tuple[str, int]]) -> None:
    """Header row + column widths. `columns` is [(title, width), ...].

    Bold on a shaded fill and frozen below, because these sheets are wide and a
    customer scrolls them. Widths are set explicitly: a description column at
    Excel's default 8.43 characters shows "CSG 10-3..." and the reason column
    would show nothing usable at all.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.append([title for title, _w in columns])
    fill = PatternFill("solid", fgColor="E9EDF4")
    for idx, (_title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"


def _note_row(sheet, columns: int, text: str) -> None:
    """A single explanatory row across an otherwise empty sheet.

    An empty sheet under a header row is indistinguishable from a bug -- the
    reader cannot tell "nothing to report" from "the export failed halfway". So
    every sheet with no data rows says so in prose, in the first cell.
    """
    sheet.append([text] + [None] * max(0, columns - 1))


# ---------------------------------------------------------------------------
# Tab 1 -- Summary
# ---------------------------------------------------------------------------

SUMMARY_COLUMNS: list[tuple[str, int]] = [
    ("Action", 24),
    ("Product", 34),
    ("Product ID", 14),
    ("Net Shortfall", 14),
    ("Demand", 14),
    ("Drawn Customer-owned", 14),
    ("Drawn Company", 14),
    ("Drawn Substitute", 14),
    ("Unit", 8),
    ("ROS Date", 12),
    ("Required Ship Date", 14),
    ("Recommended Order Date", 15),
    ("Lead Time (months)", 12),
    ("Lead Time Basis", 46),
    ("Unrecoverable", 13),
    ("Demand Lines", 12),
    ("Reason", 90),
]


_SUMMARY_INDEX = {name: i for i, (name, _w) in enumerate(SUMMARY_COLUMNS)}


def _summary_row(rec: MrpRecommendation) -> list:
    months, basis = _lead_time_cell(rec)
    return [
        _ESCALATE if rec.unrecoverable else _ORDERABLE,
        _product_label(rec.product_id, rec.product_description),
        rec.product_id,
        _qty(rec.quantity),
        _qty(rec.demand_quantity),
        _qty(rec.drawn_customer_owned),
        _qty(rec.drawn_company),
        _qty(rec.drawn_substitute),
        _unit(rec.unit_of_measure),
        _day(rec.ros_date),
        _day(rec.required_ship_date),
        _day(rec.recommended_order_date),
        months,
        basis,
        "YES" if rec.unrecoverable else "no",
        len(rec.demand_line_ids),
        rec.reason,
    ]


def _write_summary(sheet, recommendations: list[MrpRecommendation], scope: str) -> None:
    """Tab 1: the recommendation list, orderable rows first then unrecoverable.

    THE TWO GROUPS ARE KEPT VISIBLY APART, matching `MrpSummary.tsx`, which
    renders them as two separate tables under two separate headings. They are
    two different ACTIONS -- place a mill order, versus escalate because no mill
    order can land in time -- and a customer-facing sheet that interleaved them
    would invite somebody to "just order" a line the engine has already said
    cannot be ordered into position. Three things keep them apart here: the sort
    (all orderable rows, then all unrecoverable), a banner row naming each group
    and its count, and a leading `Action` column so the distinction survives
    re-sorting the sheet by any other column -- which a reader will do.
    """
    from openpyxl.styles import Font, PatternFill

    _write_header(sheet, SUMMARY_COLUMNS)
    width = len(SUMMARY_COLUMNS)

    if not recommendations:
        _note_row(
            sheet,
            width,
            "NOTHING TO REPORT -- no MRP recommendations for this scope "
            f"({scope}). Every demand line the coverage engine evaluated came "
            "back covered, covered via a substitute, or pending customer "
            "approval, so there is no mill order to recommend. This is a real "
            "result, not an export failure: MRP only recommends procurement "
            "for demand left UNCOVERED or UNRECOVERABLE, or for demand coverage "
            "has never evaluated at all.",
        )
        return

    orderable = [r for r in recommendations if not r.unrecoverable]
    escalate = [r for r in recommendations if r.unrecoverable]

    banner_fill = PatternFill("solid", fgColor="DDF0DD")
    escalate_fill = PatternFill("solid", fgColor="FBE3E3")

    for group, rows, note, fill in (
        (
            f"ORDERABLE ({len(orderable)})",
            orderable,
            "Lead time still fits. Place the mill order by the recommended "
            "order date.",
            banner_fill,
        ),
        (
            f"UNRECOVERABLE - ESCALATE ({len(escalate)})",
            escalate,
            "The recommended order date has already passed: these cannot be met "
            "by mill order even if ordered today. Rescope ROS, borrow, or "
            "source externally.",
            escalate_fill,
        ),
    ):
        sheet.append([group] + [None] * (width - 2) + [note])
        banner = sheet.cell(row=sheet.max_row, column=1)
        banner.font = Font(bold=True)
        for col in range(1, width + 1):
            sheet.cell(row=sheet.max_row, column=col).fill = fill
        if not rows:
            _note_row(
                sheet,
                width,
                "    (none in this group -- nothing to report here)",
            )
            continue
        for rec in rows:
            sheet.append(_summary_row(rec))
            # Resolved by NAME, not a hard-coded index: the F04 breakdown
            # columns shifted every date column to the right once already.
            for name in ("ROS Date", "Required Ship Date", "Recommended Order Date"):
                col = _SUMMARY_INDEX[name] + 1
                sheet.cell(row=sheet.max_row, column=col).number_format = "yyyy-mm-dd"


# ---------------------------------------------------------------------------
# Tabs 2..n -- supporting detail, product as a column
# ---------------------------------------------------------------------------

def _item_sheet_name(index: int, analysis: ByItemAnalysis) -> str:
    """"D01 CSG 9-5-8 53.50 P110..." -- unique by the ordinal prefix, legal by
    sanitisation, and within Excel's 31-character cap by truncation of the
    best-effort description that follows the prefix."""
    prefix = f"D{index + 1:02d} "
    label = analysis.product_description or analysis.product_id
    cleaned = "".join(
        "-" if ch in EXCEL_SHEET_NAME_ILLEGAL else ch for ch in label
    ).replace('"', "")
    return (prefix + cleaned)[:EXCEL_SHEET_NAME_MAX].rstrip()


# 2026-08-12 rework: the ledger mirrors the By Item screen's full-plan table.
# Incoming is split by CERTAINTY -- "On Order" is real promised POs, "Suggested
# Order" is this platform's recommendation (not placed) -- and the two carry
# different fills below, so a suggestion never wears a promise's colour.
ITEM_LEDGER_COLUMNS: list[tuple[str, int]] = [
    ("Month", 10),
    ("Opening", 12),
    ("Opening (Customer)", 13),
    ("Opening (Company)", 13),
    ("In: On Order", 13),
    ("In: Suggested Order", 15),
    ("Out: Demand", 12),
    ("of which Contingency", 14),
    ("Ending", 12),
    ("Ending (Customer)", 13),
    ("Ending (Company)", 13),
    ("Runout?", 10),
]
_LEDGER_QTY_TITLES = {
    "Opening", "Opening (Customer)", "Opening (Company)",
    "In: On Order", "In: Suggested Order",
    "Out: Demand", "of which Contingency",
    "Ending", "Ending (Customer)", "Ending (Company)",
}

ITEM_LINE_COLUMNS: list[tuple[str, int]] = [
    ("ROS Date", 12),
    ("Well", 22),
    ("Customer", 20),
    ("Profile", 12),
    ("Quantity", 12),
    ("Unit", 8),
    ("Coverage", 20),
    ("Line's Own Product", 34),
    ("Coverage Reason", 80),
]


def _item_section_banner(sheet, width: int, text: str) -> None:
    """A shaded full-width banner row naming the block that follows. The two
    stacked tables on a product tab confused readers when nothing separated
    or named them; the banner makes each block self-describing."""
    from openpyxl.styles import Font, PatternFill

    sheet.append([text])
    fill = PatternFill("solid", fgColor="DCE4EE")
    for col in range(1, width + 1):
        cell = sheet.cell(row=sheet.max_row, column=col)
        cell.fill = fill
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)


def _item_header_row(sheet, columns: list[tuple[str, int]]) -> None:
    """Bold header on a shaded fill, mid-sheet (unlike _write_header, which
    owns row 1 and the freeze pane). Column widths take the max across both
    blocks since they share letters."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.append([title for title, _w in columns])
    fill = PatternFill("solid", fgColor="E9EDF4")
    for idx, (_title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=sheet.max_row, column=idx)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        letter = get_column_letter(idx)
        if (sheet.column_dimensions[letter].width or 0) < width:
            sheet.column_dimensions[letter].width = width


def _write_item_sheet(sheet, analysis: ByItemAnalysis) -> None:
    """One product's By Item view: a title row, then the monthly ledger block,
    then the demand lines charged to it -- the same two tables the By Item
    screen shows, in the same order, each under its own named banner so the
    workbook and the screen read as one document."""
    from openpyxl.styles import Font, PatternFill

    unit = _unit(analysis.unit_of_measure)

    sheet.append([_product_label(analysis.product_id, analysis.product_description)])
    sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
    sheet.append([])

    _item_section_banner(
        sheet, len(ITEM_LEDGER_COLUMNS),
        "MONTHLY LEDGER -- opening + on-order + suggested order - demand = "
        "ending (mirrors the By Item screen). 'On Order' is promised POs; "
        "'Suggested Order' is this platform's recommendation, NOT placed.",
    )
    _item_header_row(sheet, [
        (f"{t} ({unit})" if t in _LEDGER_QTY_TITLES else t, w)
        for t, w in ITEM_LEDGER_COLUMNS
    ])
    sheet.freeze_panes = f"A{sheet.max_row + 1}"
    runout_fill = PatternFill("solid", fgColor="FBE3E3")
    # Certainty colours, matching the screen: promised POs in blue, the
    # unplaced suggestion in amber.
    on_order_fill = PatternFill("solid", fgColor="E3EBF6")
    suggested_fill = PatternFill("solid", fgColor="F7F1E4")
    on_order_col = 1 + [t for t, _w in ITEM_LEDGER_COLUMNS].index("In: On Order")
    suggested_col = (
        1 + [t for t, _w in ITEM_LEDGER_COLUMNS].index("In: Suggested Order")
    )
    ledger = analysis.ledger
    for point in ledger:
        sheet.append([
            point.month,
            _qty(point.opening_balance),
            _qty(point.opening_customer_owned),
            _qty(point.opening_company),
            _qty(point.incoming_on_order) if point.incoming_on_order else "",
            _qty(point.incoming_recommended) if point.incoming_recommended else "",
            _qty(point.demand),
            _qty(point.demand_contingency),
            _qty(point.closing_balance),
            _qty(point.closing_customer_owned),
            _qty(point.closing_company),
            "RUNOUT" if point.month == analysis.ledger_runout_month else "",
        ])
        row_idx = sheet.max_row
        if point.incoming_on_order:
            sheet.cell(row=row_idx, column=on_order_col).fill = on_order_fill
        if point.incoming_recommended:
            sheet.cell(row=row_idx, column=suggested_col).fill = suggested_fill
        if point.month == analysis.ledger_runout_month:
            for col in range(1, len(ITEM_LEDGER_COLUMNS) + 1):
                sheet.cell(row=row_idx, column=col).fill = runout_fill
    if analysis.ledger_undated_on_order > 0:
        _note_row(
            sheet, len(ITEM_LEDGER_COLUMNS),
            f"A further {_qty(analysis.ledger_undated_on_order)} {unit} is on "
            "order WITHOUT a promised arrival date -- no month to land in, so "
            "it appears in no row above.",
        )
    if not ledger:
        _note_row(sheet, len(ITEM_LEDGER_COLUMNS),
                  "NOTHING TO REPORT -- no runout series for this product.")

    sheet.append([])
    _item_section_banner(
        sheet, len(ITEM_LINE_COLUMNS),
        "DEMAND LINES charged to this product",
    )
    _item_header_row(sheet, ITEM_LINE_COLUMNS)
    for line in analysis.demand_lines:
        sheet.append([
            _day(line.ros_date),
            line.well_name,
            line.customer_name or "",
            line.profile,
            _qty(line.quantity),
            _unit(line.unit_of_measure),
            line.coverage_status or "not evaluated",
            _product_label(line.product_id, line.product_description),
            line.coverage_reason or "",
        ])
        sheet.cell(row=sheet.max_row, column=1).number_format = "yyyy-mm-dd"
    if not analysis.demand_lines:
        _note_row(sheet, len(ITEM_LINE_COLUMNS),
                  "No demand lines are charged to this product.")


INVENTORY_COLUMNS: list[tuple[str, int]] = [
    ("Product", 34),
    ("Product ID", 14),
    ("Unit", 8),
    ("On Hand (all Business Units)", 16),
    ("Assigned (carve-out of On Hand)", 16),
    ("Assigned Source", 15),
    ("Customer-Owned (all customers)", 16),
    ("Customer-Owned Source", 18),
    ("On Order", 14),
    ("On Order Source", 15),
    ("On Order Earliest Arrival", 15),
    ("On Order Latest Arrival", 15),
    ("On Order Undated", 14),
    ("Oracle Feed Live", 13),
    ("Runout Month", 12),
    ("Runout Month (if recommended order placed)", 20),
    ("Notes", 90),
]


def _write_inventory(sheet, analyses: list[ByItemAnalysis]) -> None:
    """One row per recommended product: the supply side of the justification.

    ON HAND IS AN ALL-BUSINESS-UNIT TOTAL and the header says so. It is a
    PROCUREMENT figure, not a coverage figure, and must never be read as "what
    is available to customer X" -- coverage is computed per Business Unit
    through an entirely different path.

    CUSTOMER-OWNED IS REPORTED BESIDE ON-HAND, NOT ADDED TO IT. That steel is
    not ours: it is drawn against that customer's own demand first, so it
    genuinely reduces what must be ordered, but it can never be promised to
    anybody else. Adding the two into one "available" figure is the mistake this
    layout exists to prevent.

    UNKNOWN vs ZERO is preserved cell by cell. `on_order` and `customer_owned`
    are nullable, null means no row exists anywhere, and null is written as
    `NO_DATA`. `oracle_integrated` stays whatever the engine says (False today):
    a seeded projection table is not a live feed, and the `*_source` columns say
    "synthetic" beside it.
    """
    _write_header(sheet, INVENTORY_COLUMNS)
    if not analyses:
        _note_row(
            sheet,
            len(INVENTORY_COLUMNS),
            "NOTHING TO REPORT -- no inventory positions to detail, because the "
            "summary sheet has no recommendations.",
        )
        return
    for analysis in analyses:
        inv = analysis.inventory
        notes = [
            "On Hand is the total across EVERY Business Unit -- a procurement "
            "figure, not a coverage figure.",
            "Customer-Owned is NOT part of On Hand and cannot be redirected to "
            "another customer.",
        ]
        if inv.on_order is None:
            notes.append(
                "On Order is UNKNOWN: no purchase-order row exists for this "
                "product in any Business Unit. It is deliberately not reported "
                "as 0."
            )
        if inv.customer_owned is None:
            notes.append(
                "Customer-Owned is UNKNOWN: no customer has uploaded a position "
                "for this product. It is deliberately not reported as 0."
            )
        if not inv.oracle_integrated:
            notes.append(
                "The Oracle inventory feed is NOT live; the assigned/on-order "
                "figures come from the local projection (see the source "
                "columns)."
            )
        sheet.append(
            [
                _product_label(analysis.product_id, analysis.product_description),
                analysis.product_id,
                _unit(analysis.unit_of_measure),
                _qty(inv.on_hand),
                _qty(inv.assigned),
                inv.assigned_source,
                _qty(inv.customer_owned),
                inv.customer_owned_source,
                _qty(inv.on_order),
                inv.on_order_source,
                _day(inv.on_order_earliest_arrival),
                _day(inv.on_order_latest_arrival),
                _qty(inv.on_order_undated),
                "yes" if inv.oracle_integrated else "no",
                analysis.runout_month or "does not run out in the projected window",
                analysis.runout_month_with_recommended_order
                or "does not run out in the projected window",
                " ".join(notes),
            ]
        )
        for col in (11, 12):
            sheet.cell(row=sheet.max_row, column=col).number_format = "yyyy-mm-dd"


RUNOUT_COLUMNS: list[tuple[str, int]] = [
    ("Product", 34),
    ("Product ID", 14),
    ("Series", 34),
    ("Month", 10),
    ("Opening Balance", 16),
    ("Demand", 14),
    ("Closing Balance", 16),
    ("Unit", 8),
    ("Runout", 10),
]

#: The two series are named, never merged. A hypothetical assumption must not be
#: indistinguishable from the measured baseline.
SERIES_BASELINE = "Baseline (no new order)"
SERIES_WITH_ORDER = "If recommended order is placed"


def _write_runout(sheet, analyses: list[ByItemAnalysis]) -> None:
    """The monthly balance projection, both series, one row per month.

    Two series per product, labelled in the `Series` column and NEVER merged
    into one: `runout` is the measured baseline, and
    `runout_with_recommended_order` assumes this product's own RECOVERABLE
    recommended order is placed and arrives. Collapsing them into a single curve
    would make an assumption look like a fact -- the same reason On Order is
    reported beside On Hand rather than inside it. The second series is
    deliberately absent for products with no recoverable recommendation of their
    own to inject.
    """
    _write_header(sheet, RUNOUT_COLUMNS)
    written = 0
    for analysis in analyses:
        label = _product_label(analysis.product_id, analysis.product_description)
        for series_name, points, runout_month in (
            (SERIES_BASELINE, analysis.runout, analysis.runout_month),
            (
                SERIES_WITH_ORDER,
                analysis.runout_with_recommended_order,
                analysis.runout_month_with_recommended_order,
            ),
        ):
            for point in points:
                sheet.append(
                    [
                        label,
                        analysis.product_id,
                        series_name,
                        point.month,
                        _qty(point.opening_balance),
                        _qty(point.demand),
                        _qty(point.closing_balance),
                        _unit(point.unit_of_measure),
                        "RUNOUT" if point.month == runout_month else "",
                    ]
                )
                written += 1
    if not written:
        _note_row(
            sheet,
            len(RUNOUT_COLUMNS),
            "NOTHING TO REPORT -- no runout projection to detail, because the "
            "summary sheet has no recommendations.",
        )


LEAD_TIME_COLUMNS: list[tuple[str, int]] = [
    ("Product", 34),
    ("Product ID", 14),
    ("Modelled", 10),
    ("Total Lead Time (months)", 14),
    ("Dimension", 16),
    ("Matched Value (component row)", 24),
    ("Product's Value", 24),
    ("Months", 9),
    ("Shared (wildcard match)", 12),
    ("Notes", 90),
]


def _write_lead_time(sheet, analyses: list[ByItemAnalysis]) -> None:
    """The four attribute dimensions each order date is built from.

    Present for every recommended product, including one whose model is
    INCOMPLETE -- in which case the total is written as NOT MODELLED with the
    missing dimensions named, rather than as the scalar 0 the engine reports.
    "0 months" reads as instant delivery, so printing it would make the
    least-known product look like the safest line in the book. "Configure OD/WT"
    is actionable; "0" is actively misleading.
    """
    _write_header(sheet, LEAD_TIME_COLUMNS)
    if not analyses:
        _note_row(
            sheet,
            len(LEAD_TIME_COLUMNS),
            "NOTHING TO REPORT -- no lead-time breakdowns to detail, because the "
            "summary sheet has no recommendations.",
        )
        return
    for analysis in analyses:
        label = _product_label(analysis.product_id, analysis.product_description)
        lead = analysis.lead_time
        total, basis = _lead_time_cell(analysis)
        if lead is None or not lead.components:
            missing = (
                ", ".join(lead.missing_dimensions)
                if lead is not None and lead.missing_dimensions
                else "any dimension"
            )
            sheet.append(
                [
                    label,
                    analysis.product_id,
                    "no",
                    NOT_MODELLED,
                    "(none matched)",
                    "",
                    "",
                    "",
                    "",
                    f"LEAD TIME NOT MODELLED: missing lead-time components for "
                    f"{missing}. No order date derived from it can be trusted. "
                    f"{basis}",
                ]
            )
            continue
        for component in lead.components:
            sheet.append(
                [
                    label,
                    analysis.product_id,
                    "yes" if lead.modelled else "no",
                    total,
                    component.dimension,
                    component.attribute_value,
                    component.matched_on,
                    component.months,
                    "yes" if component.shared else "no",
                    basis if lead.modelled else f"LEAD TIME NOT MODELLED: {basis}",
                ]
            )


NOTES_COLUMNS: list[tuple[str, int]] = [("How to read this workbook", 130)]


def _write_notes(
    sheet,
    scope: str,
    generated_for: date,
    export: dict,
    unavailable: list[tuple[str, str]],
) -> None:
    _write_header(sheet, NOTES_COLUMNS)
    notes = [
        f"Scope: {scope}.",
        f"Generated for planning date {generated_for.isoformat()}. Every figure "
        "in this workbook was computed at that one moment, by the same MRP "
        "engine that produces the MRP Summary and By Item screens -- not "
        "recalculated for this file. The numbers here and the numbers on screen "
        "are the same numbers.",
        f"'{SHEET_SUMMARY}' is what to do: {export['orderable']} orderable "
        f"recommendation(s) and {export['unrecoverable']} unrecoverable one(s). "
        "The two groups are separated on purpose -- they are two different "
        "actions. An ORDERABLE row is a mill order to place by the recommended "
        "order date. An UNRECOVERABLE row cannot be met by mill order even if "
        "ordered today and must be escalated: rescope the ROS date, borrow, or "
        "source externally.",
        "Each 'D01'..'Dnn' sheet is ONE recommended product -- the modern "
        "equivalent of the source workbook's 'By Item' sheet. It holds two "
        "blocks under named banners: the MONTHLY LEDGER (projected month-end "
        "stock; the runout month is shaded red) and the DEMAND LINES charged "
        "to that product. The 'Dnn' prefix keeps tab names unique within "
        "Excel's 31-character cap.",
        f"'{SHEET_INVENTORY}', '{SHEET_RUNOUT}' and '{SHEET_LEAD_TIME}' are "
        "the cross-product justification behind the summary rows; each carries "
        "a Product column instead of being split per tab.",
        f"UNKNOWN IS NOT ZERO. A cell reading '{NO_DATA}' means no source row "
        "exists for that figure at all. It is never written as 0, because a "
        "measured zero ('nothing is on order') and an absent measurement ('we "
        "do not know what is on order') are different facts and only one of them "
        "is safe to plan against.",
        f"A lead time of '{NOT_MODELLED}' means the attribute components needed "
        "to total it are not configured. That is unknown, not instant: the "
        "missing dimensions are named, and any order date shown for such a "
        "product is indicative only.",
        "Every quantity is labelled with its unit -- in the column header where "
        "a sheet's quantities all belong to one product, and in an adjacent "
        "'Unit' column where they do not. In a product sheet's DEMAND LINES "
        "block the Unit column is the DEMAND LINE's own product's unit, which "
        "can differ from the charged product's when coverage satisfied the "
        "line with a substitute.",
        "'On Hand' is the total across EVERY Business Unit, because MRP is a "
        "system-wide procurement view. It is not a coverage figure and must not "
        "be read as what is available to any one customer. 'Customer-Owned' is "
        "reported separately and is NOT ours: it will be consumed against that "
        "customer's own demand first, so it reduces what must be ordered, but it "
        "can never be promised to a third party.",
        f"On '{SHEET_RUNOUT}' the two series are labelled and never merged. "
        f"'{SERIES_BASELINE}' is the measured projection. "
        f"'{SERIES_WITH_ORDER}' additionally assumes this product's own "
        "RECOVERABLE recommended order is placed and arrives; the unrecoverable "
        "portion is deliberately NOT injected, because the engine has already "
        "determined it cannot land in time.",
        "This file is a read-only report. Editing it changes nothing in the "
        "platform -- unlike the Demand Import template, it is not an upload "
        "contract.",
    ]
    for note in notes:
        sheet.append([note])
    if unavailable:
        sheet.append([""])
        sheet.append(
            [
                "SUPPORTING DETAIL UNAVAILABLE for the product(s) below. They "
                "still appear on the summary sheet -- the recommendation itself "
                "is sound -- but their By Item justification could not be built, "
                "and the reason is stated rather than the sheets quietly "
                "omitting them:"
            ]
        )
        for label, reason in unavailable:
            sheet.append([f"    {label}: {reason}"])


# ---------------------------------------------------------------------------
# Assembly
# ---------------------------------------------------------------------------


def _analyses(
    db: Session,
    recommendations: list[MrpRecommendation],
    today: date | None,
) -> tuple[list[ByItemAnalysis], list[tuple[str, str]]]:
    """`by_item` for each DISTINCT recommended product, in summary-sheet order.

    Distinct: a product can carry two recommendation rows (a recoverable one and
    an unrecoverable one -- see `app.engines.mrp._recommendations_for_lines`),
    and its By Item justification is one analysis covering both. Duplicating it
    would double every runout row.

    `by_item` RAISES for a product with no InventoryOnHand row in any Business
    Unit, on purpose: unknown is not zero and a runout curve opening at a
    fabricated 0 is a confident wrong number. That refusal is honoured rather
    than worked around -- the product keeps its summary row (the recommendation
    needs no inventory figure) and the reason is reported on the Notes sheet.
    Letting the exception escape would mean one unseeded product makes the whole
    workbook 500, which would hide every other recommendation in it.
    """
    analyses: list[ByItemAnalysis] = []
    unavailable: list[tuple[str, str]] = []
    seen: set[str] = set()
    for rec in recommendations:
        if rec.product_id in seen:
            continue
        seen.add(rec.product_id)
        try:
            analyses.append(by_item(db, rec.product_id, today=today))
        except Exception as exc:  # noqa: BLE001 -- reported, never swallowed
            unavailable.append(
                (
                    _product_label(rec.product_id, rec.product_description),
                    f"{type(exc).__name__}: {exc}",
                )
            )
    return analyses, unavailable


def _safe(text: str, fallback: str) -> str:
    return (
        "".join(ch if (ch.isalnum() or ch in "-_") else "-" for ch in text).strip("-")
        or fallback
    )


def build_mrp_export(
    db: Session,
    customer=None,
    business_unit_id: str | None = None,
    today: date | None = None,
) -> MrpExport:
    """The MRP workbook: summary on tab 1, supporting detail on tabs 2 onward.

    `customer` is an optional `Customer` model instance, mirroring
    `mrp_summary`'s optional `customer_id` -- MRP is system-wide by default and
    this export must not be narrower than the screen it exports.

    `today` is resolved ONCE here and threaded into both `mrp_summary` and every
    `by_item` call, so the summary sheet and its supporting sheets cannot
    describe two different moments. Left to default independently, a request
    crossing midnight would produce a workbook whose recoverability verdicts and
    runout curves disagreed.

    Reads only. Generating a report is a read of the plan.
    """
    from openpyxl import Workbook

    today = today or date.today()
    scope = (
        f"customer {customer.name!r}"
        if customer is not None
        else "ALL customers (system-wide, matching the MRP Summary screen's default)"
    )

    recommendations = mrp_summary(
        db,
        customer_id=(customer.id if customer is not None else None),
        business_unit_id=business_unit_id,
        today=today
    )
    analyses, unavailable = _analyses(db, recommendations, today)

    workbook = Workbook()
    summary = workbook.active
    summary.title = SHEET_SUMMARY
    _write_summary(summary, recommendations, scope)
    item_names = [
        _item_sheet_name(i, analysis) for i, analysis in enumerate(analyses)
    ]
    assert len(set(item_names)) == len(item_names), item_names
    for name, analysis in zip(item_names, analyses):
        _write_item_sheet(workbook.create_sheet(name), analysis)
    _write_inventory(workbook.create_sheet(SHEET_INVENTORY), analyses)
    _write_runout(workbook.create_sheet(SHEET_RUNOUT), analyses)
    _write_lead_time(workbook.create_sheet(SHEET_LEAD_TIME), analyses)
    orderable = sum(1 for r in recommendations if not r.unrecoverable)
    _write_notes(
        workbook.create_sheet(SHEET_NOTES),
        scope,
        today,
        {
            "orderable": orderable,
            "unrecoverable": len(recommendations) - orderable,
        },
        unavailable,
    )

    sheet_names = tuple(workbook.sheetnames)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    who = _safe(customer.name, "customer") if customer is not None else "all-customers"
    return MrpExport(
        content=buffer.getvalue(),
        filename=f"mrp-{who}-{today.isoformat()}.xlsx",
        recommendation_count=len(recommendations),
        orderable_count=orderable,
        unrecoverable_count=len(recommendations) - orderable,
        customer_name=customer.name if customer is not None else None,
        sheet_names=sheet_names,
        unavailable_products=tuple(unavailable),
    )

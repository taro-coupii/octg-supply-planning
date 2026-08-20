# C-01R resolved (stage 7): protected via app.main's include_router(dependencies=[...]).
"""spec §API rows 1-4: MRP summary/detail/export, MOR order-requirements.

API-layer composition (spec §MRPエンジン receipts_recommended note): /mrp/summary
and /mrp/items call mor_rows FIRST and feed its requirements into mrp_rows as
`recommended_by_product` — the engines themselves stay independent; only this
layer wires MOR's recommendations into MRP's receipts_recommended column.
"""

from __future__ import annotations

import io
from collections import defaultdict

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.db import get_db
from app.engines.coverage import _scope, scoped_lines
from app.engines.mor import MorRow, mor_rows
from app.engines.mrp import ProductLedger, mrp_rows
from app.models import BusinessUnit, Customer, Product

router = APIRouter(prefix="/mrp")

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FONT = Font(bold=True)


class ScopeOut(BaseModel):
    statuses: list[str]
    profiles: list[str]


def _scope_out(db: Session) -> ScopeOut:
    statuses, profiles = _scope(db)
    return ScopeOut(statuses=sorted(statuses), profiles=sorted(profiles))


class BalanceOut(BaseModel):
    company: float
    owned: float


class MonthOut(BaseModel):
    month: str
    opening: BalanceOut
    receipts_booked: float
    receipts_recommended: float
    issues: float
    closing: BalanceOut


class RunoutOut(BaseModel):
    baseline: str | None
    with_recommended: str | None
    on_order: str | None


class ProductSummaryOut(BaseModel):
    product: str
    unit: str
    business_unit_id: str
    business_unit_name: str
    opening: BalanceOut
    runout_months: RunoutOut
    months: list[MonthOut]
    on_order_undated: float


class MrpSummaryOut(BaseModel):
    rows: list[ProductSummaryOut]
    scope: ScopeOut


def _recommended_by_bu_product(mor_result) -> dict[tuple[str, str], dict[str, float]]:
    """Sum MOR requirement qty by ((bu_id, product_id), need_month). Keying by
    (bu, product) — not product alone — keeps requirements from different BUs
    from merging into each other's ledger (a shared product across BUs must
    never double-count or leak one BU's requirement into another's)."""
    out: dict[tuple[str, str], dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for row in mor_result.rows:
        for req in row.requirements:
            out[(row.bu_id, row.product_id)][req.need_month] += req.qty
    return out


def _bu_name_map(db: Session) -> dict[str, str]:
    return {bu.id: bu.name for bu in db.query(BusinessUnit).all()}


def _balance_out(b) -> BalanceOut:
    return BalanceOut(company=b.company, owned=b.owned)


def _months_out(ledger: ProductLedger) -> list[MonthOut]:
    return [
        MonthOut(
            month=m.month,
            opening=_balance_out(m.opening),
            receipts_booked=m.receipts_booked,
            receipts_recommended=m.receipts_recommended,
            issues=m.issues,
            closing=_balance_out(m.closing),
        )
        for m in ledger.months
    ]


def _summary_out(ledger: ProductLedger, bu_names: dict[str, str]) -> ProductSummaryOut:
    return ProductSummaryOut(
        product=ledger.product_id,
        unit=ledger.unit.value,
        business_unit_id=ledger.bu_id,
        business_unit_name=bu_names.get(ledger.bu_id, ledger.bu_id),
        opening=_balance_out(ledger.opening),
        runout_months=RunoutOut(**ledger.runout_months),
        months=_months_out(ledger),
        on_order_undated=ledger.on_order_undated,
    )


@router.get("/summary", response_model=MrpSummaryOut)
def get_mrp_summary(horizon: int | None = None, db: Session = Depends(get_db)):
    mor_result = mor_rows(db, horizon=horizon)
    recommended = _recommended_by_bu_product(mor_result)
    ledgers = mrp_rows(db, horizon=horizon, recommended_by_product=recommended)
    bu_names = _bu_name_map(db)
    return MrpSummaryOut(rows=[_summary_out(l, bu_names) for l in ledgers], scope=_scope_out(db))


class DemandLineOut(BaseModel):
    id: str
    customer_id: str
    well_id: str
    quantity: float
    unit: str
    ros_date: str
    profile: str


class ProductItemOut(ProductSummaryOut):
    lines: list[DemandLineOut]


class MrpItemOut(BaseModel):
    row: ProductItemOut
    scope: ScopeOut


def _matching_ledger(ledgers: list[ProductLedger], product_id: str, bu_id: str | None) -> ProductLedger | None:
    candidates = [l for l in ledgers if l.product_id == product_id]
    if bu_id is not None:
        candidates = [l for l in candidates if l.bu_id == bu_id]
    return candidates[0] if candidates else None


@router.get("/items/{product_id}", response_model=MrpItemOut)
def get_mrp_item(product_id: str, horizon: int | None = None, bu_id: str | None = None, db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if product is None:
        raise HTTPException(status_code=404, detail="Product not found")

    mor_result = mor_rows(db, horizon=horizon)
    recommended = _recommended_by_bu_product(mor_result)
    ledgers = mrp_rows(db, horizon=horizon, recommended_by_product=recommended)
    ledger = _matching_ledger(ledgers, product_id, bu_id)
    if ledger is None:
        raise HTTPException(status_code=404, detail="No MRP ledger for this product")

    statuses, profiles = _scope(db)
    lines_out: list[DemandLineOut] = []
    customers = db.query(Customer).filter(Customer.business_unit_id == ledger.bu_id).all()
    for cust in customers:
        for line in scoped_lines(db, cust.id, statuses, profiles):
            if line.product_id != product_id:
                continue
            lines_out.append(
                DemandLineOut(
                    id=line.id,
                    customer_id=cust.id,
                    well_id=line.well_id,
                    quantity=line.quantity,
                    unit=line.unit.value,
                    ros_date=line.ros_date.isoformat(),
                    profile=line.profile.value,
                )
            )
    lines_out.sort(key=lambda l: l.ros_date)

    summary = _summary_out(ledger, _bu_name_map(db))
    item = ProductItemOut(**summary.model_dump(), lines=lines_out)
    return MrpItemOut(row=item, scope=_scope_out(db))


class RequirementOut(BaseModel):
    need_month: str
    qty: float
    unit: str
    ex_mill_month: str
    overdue: bool


class StripPointOut(BaseModel):
    month: str
    closing_balance: float


class MarkersOut(BaseModel):
    order_deadline: str | None
    physical_runout: str | None
    safety_breach: str | None


class MorProductOut(BaseModel):
    product: str
    unit: str
    business_unit_id: str
    business_unit_name: str
    strip: list[StripPointOut]
    markers: MarkersOut
    requirements: list[RequirementOut]
    safety_stock: float | None
    unavailable_reason: str | None = None


class MorSummaryOut(BaseModel):
    rows: list[MorProductOut]
    scope: ScopeOut
    unavailable_reason: str | None = None


def _mor_row_out(row: MorRow, bu_names: dict[str, str]) -> MorProductOut:
    return MorProductOut(
        product=row.product_id,
        unit=row.unit.value,
        business_unit_id=row.bu_id,
        business_unit_name=bu_names.get(row.bu_id, row.bu_id),
        strip=[StripPointOut(**p) for p in row.strip],
        markers=MarkersOut(**row.markers),
        requirements=[
            RequirementOut(
                need_month=r.need_month, qty=r.qty, unit=r.unit.value, ex_mill_month=r.ex_mill_month, overdue=r.overdue
            )
            for r in row.requirements
        ],
        safety_stock=row.safety_stock,
    )


@router.get("/order-requirements", response_model=MorSummaryOut)
def get_order_requirements(horizon: int | None = None, customer_id: str | None = None, db: Session = Depends(get_db)):
    result = mor_rows(db, horizon=horizon, customer_id=customer_id)
    bu_names = _bu_name_map(db)
    return MorSummaryOut(
        rows=[_mor_row_out(r, bu_names) for r in result.rows],
        scope=_scope_out(db),
        unavailable_reason=result.unavailable_reason,
    )


MONTHLY_HEADERS = [
    "Month",
    "Opening company",
    "Opening owned",
    "Receipts booked",
    "Receipts recommended",
    "Issues",
    "Closing company",
    "Closing owned",
]

DEMAND_HEADERS = ["Need month", "Well", "Customer", "Quantity", "Unit", "Profile"]


def _product_label(db: Session, product_id: str) -> str:
    p = db.get(Product, product_id)
    return p.name if p is not None else product_id


_INVALID_SHEET_CHARS = set('\\/*?:[]')


def _sheet_title(label: str, product_id: str) -> str:
    """Excel worksheet titles forbid \\/*?:[] and are capped at 31 chars
    (product names/IDs may contain any of these, e.g. '9-5/8"' sizes)."""
    cleaned = "".join(c if c not in _INVALID_SHEET_CHARS else "-" for c in label).strip()
    if not cleaned:
        cleaned = product_id
    return cleaned[:31]


def _build_workbook(db: Session, ledgers: list[ProductLedger], statuses, profiles) -> Workbook:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws.append(["Product", "Unit", "Baseline runout", "With recommended runout", "On order runout", "On order undated"])
    for cell in summary_ws[1]:
        cell.font = HEADER_FONT
    for ledger in ledgers:
        summary_ws.append(
            [
                _product_label(db, ledger.product_id),
                ledger.unit.value,
                ledger.runout_months.get("baseline"),
                ledger.runout_months.get("with_recommended"),
                ledger.runout_months.get("on_order"),
                ledger.on_order_undated,
            ]
        )

    for ledger in ledgers:
        name = _sheet_title(_product_label(db, ledger.product_id), ledger.product_id)
        if name in wb.sheetnames:
            suffix = f" {ledger.product_id[:4]}"
            name = name[: 31 - len(suffix)] + suffix
        ws = wb.create_sheet(title=name)
        ws.append([f"{_product_label(db, ledger.product_id)} ({ledger.unit.value})"])
        ws.append([])

        header_row_idx = ws.max_row + 1
        ws.append(["MONTHLY LEDGER"])
        for cell in ws[ws.max_row]:
            cell.font = HEADER_FONT
        ws.append(MONTHLY_HEADERS)
        for cell in ws[ws.max_row]:
            cell.font = HEADER_FONT

        baseline_runout = ledger.runout_months.get("baseline")
        for m in ledger.months:
            row_idx = ws.max_row + 1
            ws.append(
                [
                    m.month,
                    m.opening.company,
                    m.opening.owned,
                    m.receipts_booked,
                    m.receipts_recommended,
                    m.issues,
                    m.closing.company,
                    m.closing.owned,
                ]
            )
            if baseline_runout is not None and m.month == baseline_runout:
                ws.cell(row=row_idx, column=1).fill = RED_FILL

        ws.append([])
        ws.append(["DEMAND LINES"])
        for cell in ws[ws.max_row]:
            cell.font = HEADER_FONT
        ws.append(DEMAND_HEADERS)
        for cell in ws[ws.max_row]:
            cell.font = HEADER_FONT

        customers = db.query(Customer).filter(Customer.business_unit_id == ledger.bu_id).all()
        demand_lines = []
        for cust in customers:
            for line in scoped_lines(db, cust.id, statuses, profiles):
                if line.product_id != ledger.product_id:
                    continue
                demand_lines.append((line, cust.name))
        demand_lines.sort(key=lambda pair: pair[0].ros_date)
        for line, cust_name in demand_lines:
            ws.append(
                [
                    line.ros_date.isoformat(),
                    line.well_id,
                    cust_name,
                    line.quantity,
                    line.unit.value,
                    line.profile.value,
                ]
            )

    return wb


@router.get("/export")
def export_mrp(horizon: int | None = None, db: Session = Depends(get_db)):
    mor_result = mor_rows(db, horizon=horizon)
    recommended = _recommended_by_bu_product(mor_result)
    ledgers = mrp_rows(db, horizon=horizon, recommended_by_product=recommended)
    statuses, profiles = _scope(db)

    wb = _build_workbook(db, ledgers, statuses, profiles)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mrp_export.xlsx"},
    )

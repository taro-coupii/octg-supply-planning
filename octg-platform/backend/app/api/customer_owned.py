import io

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, ConfigDict
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import Customer, CustomerOwnedInventory, CustomerOwnedUpload, Product, UnitOfMeasure

router = APIRouter(prefix="/customer-owned-inventory")

VALID_UNITS = {u.value for u in UnitOfMeasure}

TEMPLATE_HEADER = ["Product", "Quantity", "Unit"]


class PositionOut(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: str
    product_id: str
    quantity: float
    unit: str


class CustomerOwnedOut(BaseModel):
    has_uploaded: bool
    uploaded_at: str | None
    positions: list[PositionOut]


@router.get("", response_model=CustomerOwnedOut)
def get_customer_owned(customer_id: str, db: Session = Depends(get_db)):
    latest_upload = (
        db.query(CustomerOwnedUpload)
        .filter_by(customer_id=customer_id)
        .order_by(CustomerOwnedUpload.uploaded_at.desc())
        .first()
    )
    positions = db.query(CustomerOwnedInventory).filter_by(customer_id=customer_id).all()
    return CustomerOwnedOut(
        has_uploaded=latest_upload is not None,
        uploaded_at=latest_upload.uploaded_at.isoformat() if latest_upload else None,
        positions=positions,
    )


@router.get("/template")
def get_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(TEMPLATE_HEADER)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=customer-owned-inventory-template.xlsx"},
    )


@router.post("/upload", response_model=CustomerOwnedOut)
async def upload_customer_owned(customer_id: str, file: UploadFile, db: Session = Depends(get_db)):
    if db.query(Customer).filter_by(id=customer_id).first() is None:
        raise HTTPException(status_code=404, detail="customer not found")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active

    products_by_name = {p.name: p for p in db.query(Product).all()}

    errors: list[dict] = []
    validated_rows: list[tuple] = []
    seen_product_rows: dict[str, int] = {}  # product_id -> first row it appeared on

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None for cell in row):
            continue
        product_name, quantity, unit = (list(row) + [None, None, None])[:3]

        row_errors = []
        product = products_by_name.get(product_name)
        if product is None:
            row_errors.append(f"unknown product: {product_name!r}")
        is_number = isinstance(quantity, (int, float)) and not isinstance(quantity, bool)
        if not is_number or quantity < 0:
            row_errors.append(f"invalid quantity: {quantity!r}")
        if unit not in VALID_UNITS:
            row_errors.append(f"invalid unit: {unit!r}")

        if not row_errors and product is not None:
            first_row = seen_product_rows.get(product.id)
            if first_row is not None:
                row_errors.append(
                    f"row {idx}: duplicate product {product_name!r} (also row {first_row})"
                )
            else:
                seen_product_rows[product.id] = idx

        if row_errors:
            errors.append({"row": idx, "errors": row_errors})
            continue

        validated_rows.append((product.id, quantity, unit))

    if errors:
        raise HTTPException(status_code=422, detail=errors)

    # Full replace, atomic within this request/transaction: only delete+insert
    # after every row has validated successfully (spec §不変条件3).
    db.query(CustomerOwnedInventory).filter_by(customer_id=customer_id).delete()
    for product_id, quantity, unit in validated_rows:
        db.add(
            CustomerOwnedInventory(
                customer_id=customer_id, product_id=product_id, quantity=quantity, unit=unit
            )
        )
    db.add(
        CustomerOwnedUpload(
            customer_id=customer_id, filename=file.filename or "upload.xlsx", row_count=len(validated_rows)
        )
    )
    db.commit()

    return get_customer_owned(customer_id, db)

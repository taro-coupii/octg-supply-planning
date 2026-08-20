import datetime
import io

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import (
    Customer,
    DemandImport,
    DemandImportRow,
    DemandImportStatus,
    DemandLine,
    DemandProfile,
    Product,
    UnitOfMeasure,
    Well,
)
from app.services.demand_apply import apply_import_rows

router = APIRouter(prefix="/demand")

TEMPLATE_HEADER = ["Well", "Product", "Quantity", "Unit", "ROS Date", "Profile"]
VALID_UNITS = {u.value for u in UnitOfMeasure}
VALID_PROFILES = {p.value for p in DemandProfile}


class ConflictOut(BaseModel):
    well_name: str
    existing_line_count: int
    staged_line_count: int
    existing_qty_by_unit: dict
    staged_qty_by_unit: dict


class ImportOut(BaseModel):
    id: str
    customer_id: str
    filename: str
    status: str
    conflicts: list[ConflictOut]


class ApplyOut(BaseModel):
    applied_wells: list[str]
    created_wells: list[str]


def _qty_by_unit(lines) -> dict:
    result: dict = {}
    for line in lines:
        unit = line.unit.value if hasattr(line.unit, "value") else line.unit
        result[unit] = result.get(unit, 0) + line.quantity
    return result


def _compute_conflicts(imp: DemandImport, db: Session) -> list[ConflictOut]:
    staged_rows = db.query(DemandImportRow).filter_by(import_id=imp.id).all()
    by_well: dict[str, list[DemandImportRow]] = {}
    for row in staged_rows:
        by_well.setdefault(row.well_name, []).append(row)

    conflicts = []
    for well_name, rows in by_well.items():
        well = db.query(Well).filter_by(customer_id=imp.customer_id, name=well_name).first()
        if well is None:
            continue
        existing_lines = db.query(DemandLine).filter_by(well_id=well.id).all()
        if not existing_lines:
            continue
        conflicts.append(
            ConflictOut(
                well_name=well_name,
                existing_line_count=len(existing_lines),
                staged_line_count=len(rows),
                existing_qty_by_unit=_qty_by_unit(existing_lines),
                staged_qty_by_unit=_qty_by_unit(rows),
            )
        )
    return conflicts


def _out(imp: DemandImport, db: Session) -> ImportOut:
    return ImportOut(
        id=imp.id,
        customer_id=imp.customer_id,
        filename=imp.filename,
        status=imp.status.value,
        conflicts=_compute_conflicts(imp, db),
    )


@router.get("/template")
def get_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(TEMPLATE_HEADER)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=demand-template.xlsx"},
    )


@router.post("/imports", response_model=ImportOut)
async def upload_import(customer_id: str, file: UploadFile, db: Session = Depends(get_db)):
    customer = db.get(Customer, customer_id)
    if customer is None:
        raise HTTPException(status_code=404, detail="Customer not found")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception:
        raise HTTPException(status_code=422, detail="not a valid xlsx file")
    ws = wb.active

    products_by_name = {p.name: p for p in db.query(Product).all()}

    errors: list[dict] = []
    validated_rows: list[dict] = []

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None for cell in row):
            continue
        well_name, product_name, quantity, unit, ros_date, profile = (
            list(row) + [None] * 6
        )[:6]

        row_errors = []
        if not well_name:
            row_errors.append("missing well name")
        product = products_by_name.get(product_name)
        if product is None:
            row_errors.append(f"unknown product: {product_name!r}")
        is_number = isinstance(quantity, (int, float)) and not isinstance(quantity, bool)
        if not is_number or quantity <= 0:
            row_errors.append(f"invalid quantity: {quantity!r}")
        if unit not in VALID_UNITS:
            row_errors.append(f"invalid unit: {unit!r}")
        if profile not in VALID_PROFILES:
            row_errors.append(f"invalid profile: {profile!r}")

        parsed_date = None
        if isinstance(ros_date, datetime.datetime):
            parsed_date = ros_date.date()
        elif isinstance(ros_date, datetime.date):
            parsed_date = ros_date
        elif isinstance(ros_date, str):
            try:
                parsed_date = datetime.date.fromisoformat(ros_date)
            except ValueError:
                row_errors.append(f"invalid ROS date: {ros_date!r}")
        else:
            row_errors.append(f"invalid ROS date: {ros_date!r}")

        if row_errors:
            errors.append({"row": idx, "errors": row_errors})
            continue

        validated_rows.append(
            {
                "row_no": idx,
                "well_name": well_name,
                "product_id": product.id,
                "quantity": quantity,
                "unit": unit,
                "ros_date": parsed_date,
                "profile": profile,
            }
        )

    if errors:
        raise HTTPException(status_code=422, detail=errors)

    imp = DemandImport(
        customer_id=customer_id,
        filename=file.filename or "demand.xlsx",
        status=DemandImportStatus.PENDING,
    )
    db.add(imp)
    db.flush()
    for r in validated_rows:
        db.add(
            DemandImportRow(
                import_id=imp.id,
                row_no=r["row_no"],
                well_name=r["well_name"],
                product_id=r["product_id"],
                quantity=r["quantity"],
                unit=UnitOfMeasure(r["unit"]),
                ros_date=r["ros_date"],
                profile=DemandProfile(r["profile"]),
            )
        )
    db.commit()

    return _out(imp, db)


@router.get("/imports", response_model=list[ImportOut])
def list_imports(db: Session = Depends(get_db)):
    imports = db.query(DemandImport).order_by(DemandImport.uploaded_at.desc()).all()
    return [_out(i, db) for i in imports]


@router.get("/imports/{import_id}", response_model=ImportOut)
def get_import(import_id: str, db: Session = Depends(get_db)):
    imp = db.get(DemandImport, import_id)
    if imp is None:
        raise HTTPException(status_code=404, detail="Import not found")
    return _out(imp, db)


@router.post("/imports/{import_id}/apply", response_model=ApplyOut)
def apply_import(import_id: str, db: Session = Depends(get_db)):
    imp = db.get(DemandImport, import_id)
    if imp is None:
        raise HTTPException(status_code=404, detail="Import not found")
    if imp.status != DemandImportStatus.PENDING:
        raise HTTPException(status_code=409, detail="Import is not pending")

    applied_wells, created_wells = apply_import_rows(db, imp)

    imp.status = DemandImportStatus.APPLIED
    db.commit()

    return ApplyOut(applied_wells=applied_wells, created_wells=created_wells)


@router.post("/imports/{import_id}/discard", response_model=ImportOut)
def discard_import(import_id: str, db: Session = Depends(get_db)):
    imp = db.get(DemandImport, import_id)
    if imp is None:
        raise HTTPException(status_code=404, detail="Import not found")
    if imp.status != DemandImportStatus.PENDING:
        raise HTTPException(status_code=409, detail="Import is not pending")

    imp.status = DemandImportStatus.DISCARDED
    db.commit()
    return _out(imp, db)

import datetime
import io

import openpyxl

from app.models import (
    Customer,
    DemandImport,
    DemandImportStatus,
    DemandLine,
    DemandProfile,
    DemandStatus,
    Product,
    UnitOfMeasure,
    Well,
)

TEMPLATE_HEADER = ["Well", "Product", "Quantity", "Unit", "ROS Date", "Profile"]


def _fixtures(db):
    c1 = Customer(name="C1")
    c2 = Customer(name="C2")
    p = Product(name="P1", unit_of_measure=UnitOfMeasure.MTR)
    db.add_all([c1, c2, p])
    db.commit()
    return c1, c2, p


def _xlsx_bytes(rows, header=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header or TEMPLATE_HEADER)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _upload(client, customer_id, rows):
    data = _xlsx_bytes(rows)
    return client.post(
        f"/demand/imports?customer_id={customer_id}",
        files={
            "file": (
                "demand.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_template_downloads_with_expected_header(client, db):
    resp = client.get("/demand/template")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == TEMPLATE_HEADER


def test_upload_stages_rows_and_returns_pending_with_conflicts(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(
        client,
        c1.id,
        [
            ["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"],
            ["WellB", p.name, 5, "Mtr", "2026-04-01", "Contingency"],
        ],
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "pending"
    assert body["conflicts"] == []


def test_upload_conflict_reports_existing_vs_staged(client, db):
    c1, c2, p = _fixtures(db)
    well = Well(customer_id=c1.id, name="WellA")
    db.add(well)
    db.commit()
    db.add(
        DemandLine(
            well_id=well.id,
            product_id=p.id,
            quantity=3.0,
            unit=UnitOfMeasure.MTR,
            ros_date=datetime.date(2026, 1, 1),
            profile=DemandProfile.PRIMARY,
        )
    )
    db.commit()

    resp = _upload(
        client,
        c1.id,
        [
            ["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"],
            ["WellA", p.name, 5, "Mtr", "2026-04-01", "Contingency"],
        ],
    )
    assert resp.status_code == 200
    body = resp.json()
    conflicts = body["conflicts"]
    assert len(conflicts) == 1
    conflict = conflicts[0]
    assert conflict["well_name"] == "WellA"
    assert conflict["existing_line_count"] == 1
    assert conflict["staged_line_count"] == 2
    assert conflict["existing_qty_by_unit"] == {"Mtr": 3.0}
    assert conflict["staged_qty_by_unit"] == {"Mtr": 15.0}


def test_upload_validation_failure_leaves_no_staging(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(
        client,
        c1.id,
        [
            ["WellA", "Unknown Product", 10, "Mtr", "2026-03-01", "Primary"],
            ["WellB", p.name, -1, "Mtr", "2026-04-01", "Primary"],
            ["WellC", p.name, 5, "BadUnit", "2026-04-01", "Primary"],
            ["WellD", p.name, 5, "Mtr", "not-a-date", "Primary"],
            ["WellE", p.name, 5, "Mtr", "2026-04-01", "BadProfile"],
        ],
    )
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert isinstance(detail, list)
    assert len(detail) == 5
    row_numbers = {d["row"] for d in detail}
    assert row_numbers == {2, 3, 4, 5, 6}

    assert db.query(DemandImport).count() == 0


def test_apply_full_replaces_and_creates_wells_and_revisions(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(
        client,
        c1.id,
        [
            ["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"],
        ],
    )
    import_id = resp.json()["id"]

    resp = client.post(f"/demand/imports/{import_id}/apply")
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_wells"] == ["WellA"]
    assert body["applied_wells"] == ["WellA"]

    well = db.query(Well).filter_by(customer_id=c1.id, name="WellA").one()
    assert well.demand_status == DemandStatus.PLANNED
    lines = db.query(DemandLine).filter_by(well_id=well.id).all()
    assert len(lines) == 1
    assert lines[0].quantity == 10

    imp = db.get(DemandImport, import_id)
    assert imp.status == DemandImportStatus.APPLIED


def test_apply_replaces_existing_lines_for_conflicting_well(client, db):
    c1, c2, p = _fixtures(db)
    well = Well(customer_id=c1.id, name="WellA")
    db.add(well)
    db.commit()
    db.add(
        DemandLine(
            well_id=well.id,
            product_id=p.id,
            quantity=3.0,
            unit=UnitOfMeasure.MTR,
            ros_date=datetime.date(2026, 1, 1),
            profile=DemandProfile.PRIMARY,
        )
    )
    db.commit()

    resp = _upload(
        client,
        c1.id,
        [
            ["WellA", p.name, 20, "Mtr", "2026-03-01", "Primary"],
        ],
    )
    import_id = resp.json()["id"]
    resp = client.post(f"/demand/imports/{import_id}/apply")
    assert resp.status_code == 200

    lines = db.query(DemandLine).filter_by(well_id=well.id).all()
    assert len(lines) == 1
    assert lines[0].quantity == 20


def test_apply_records_revision_no_increments(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(client, c1.id, [["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"]])
    import_id = resp.json()["id"]
    client.post(f"/demand/imports/{import_id}/apply")

    well = db.query(Well).filter_by(customer_id=c1.id, name="WellA").one()

    resp2 = _upload(client, c1.id, [["WellA", p.name, 15, "Mtr", "2026-04-01", "Primary"]])
    import_id2 = resp2.json()["id"]
    client.post(f"/demand/imports/{import_id2}/apply")

    from app.models import DemandRevision

    revisions = (
        db.query(DemandRevision).filter_by(well_id=well.id).order_by(DemandRevision.revision_no).all()
    )
    assert [r.revision_no for r in revisions] == [1, 2]
    assert all(r.source.value == "import" for r in revisions)


def test_apply_non_pending_returns_409(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(client, c1.id, [["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"]])
    import_id = resp.json()["id"]
    client.post(f"/demand/imports/{import_id}/apply")

    resp = client.post(f"/demand/imports/{import_id}/apply")
    assert resp.status_code == 409


def test_discard_pending_transitions_status_and_keeps_staging(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(client, c1.id, [["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"]])
    import_id = resp.json()["id"]

    resp = client.post(f"/demand/imports/{import_id}/discard")
    assert resp.status_code == 200
    assert resp.json()["status"] == "discarded"

    from app.models import DemandImportRow

    assert db.query(DemandImportRow).filter_by(import_id=import_id).count() == 1


def test_discard_non_pending_returns_409(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(client, c1.id, [["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"]])
    import_id = resp.json()["id"]
    client.post(f"/demand/imports/{import_id}/discard")

    resp = client.post(f"/demand/imports/{import_id}/discard")
    assert resp.status_code == 409


def test_list_and_get_import(client, db):
    c1, c2, p = _fixtures(db)
    resp = _upload(client, c1.id, [["WellA", p.name, 10, "Mtr", "2026-03-01", "Primary"]])
    import_id = resp.json()["id"]

    resp = client.get("/demand/imports")
    assert resp.status_code == 200
    assert len(resp.json()) == 1

    resp = client.get(f"/demand/imports/{import_id}")
    assert resp.status_code == 200
    assert resp.json()["id"] == import_id
    assert "conflicts" in resp.json()


def test_get_import_not_found_404(client, db):
    resp = client.get("/demand/imports/nonexistent")
    assert resp.status_code == 404


def test_upload_unknown_customer_returns_404(client, db):
    _fixtures(db)
    resp = _upload(client, "nonexistent-customer", [["WellA", "P1", 10, "Mtr", "2026-03-01", "Primary"]])
    assert resp.status_code == 404
    assert resp.json()["detail"] == "Customer not found"


def test_upload_non_xlsx_file_returns_422(client, db):
    c1, c2, p = _fixtures(db)
    resp = client.post(
        f"/demand/imports?customer_id={c1.id}",
        files={"file": ("demand.xlsx", b"this is not an xlsx file", "text/plain")},
    )
    assert resp.status_code == 422
    assert resp.json()["detail"] == "not a valid xlsx file"

    from app.models import DemandImport, DemandImportRow

    assert db.query(DemandImport).count() == 0
    assert db.query(DemandImportRow).count() == 0

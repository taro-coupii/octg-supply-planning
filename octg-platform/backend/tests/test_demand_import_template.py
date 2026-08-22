"""Template download: the CURRENT demand book, in the shape the parser reads back.

THE CORRECTNESS BAR, and it is a round trip
-------------------------------------------
`GET /demand-imports/template?customer_id=...` is not a blank column hint. It is an
export of live demand, and its whole value rests on one property:

    re-uploading it UNMODIFIED must stage every row as an exact-match revision,
    with no errors and no conflicts, changing nothing.

A planner edits the cells they mean to change; nothing else may move. That property
is not checkable by inspecting the generator -- it is a statement about the generator
and the parser agreeing -- so it is tested by actually generating a file, actually
posting it back through the real upload endpoint, and asserting the staged result
reports nothing changed. `test_generated_template_round_trips_with_no_changes` is
the load-bearing test in this module.
"""

import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.engines.coverage import apply_revision, set_well_demand_status
from app.engines.demand_import import TEMPLATE_COLUMNS
from app.main import app
from app.models import DemandLine, DemandProfile, DemandStatus, Well
from tests.phase5_fixtures import NOW, build_client


@pytest.fixture()
def client_world():
    client, session_factory, world = build_client()
    try:
        yield client, session_factory, world
    finally:
        app.dependency_overrides.clear()


def _download(client, customer_id, well_id=None):
    query = f"?customer_id={customer_id}" + (
        f"&well_id={well_id}" if well_id else ""
    )
    resp = client.get(f"/demand-imports/template{query}")
    assert resp.status_code == 200, resp.text
    return resp


def _sheet_rows(payload: bytes, sheet=0):
    wb = load_workbook(io.BytesIO(payload), data_only=True)
    ws = wb.worksheets[sheet]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    title = ws.title
    wb.close()
    return title, rows


def _upload(client, payload, filename="edited.xlsx"):
    return client.post(
        "/demand-imports",
        files={
            "file": (
                filename,
                payload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


# --------------------------------------------------------------------------
# Shape
# --------------------------------------------------------------------------


def test_template_is_a_real_xlsx_with_the_canonical_header(client_world):
    client, _sf, w = client_world
    resp = _download(client, w.acme_id)

    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment;" in resp.headers["content-disposition"]
    assert ".xlsx" in resp.headers["content-disposition"]
    # So a client can caption "5 demand lines" without opening the workbook.
    assert resp.headers["X-Demand-Row-Count"] == "5"

    title, rows = _sheet_rows(resp.content)
    assert title == "Demand"
    # The header is the CANONICAL contract, in order -- not an alias spelling that a
    # later edit to `_ALIASES` could silently break.
    assert rows[0] == TEMPLATE_COLUMNS
    assert list(TEMPLATE_COLUMNS) == [
        "well",
        "product",
        "quantity",
        "ros_date",
        "status",
        "profile",
    ]


def test_template_body_is_the_customers_current_demand_one_row_per_line(client_world):
    """Acme owns 5 demand lines across 4 wells. Every one is exported, including
    lines whose WELL is out of the default coverage scope (WELL-4 Budgeted, WELL-5
    Planned) -- the export is the demand BOOK, not the coverage answer, and a planner
    editing a Budgeted well's programme needs its rows in the file."""
    client, _sf, w = client_world
    _title, rows = _sheet_rows(_download(client, w.acme_id).content)
    body = rows[1:]
    assert len(body) == 5

    wells = sorted(r[0] for r in body)
    assert wells == ["WELL-1", "WELL-2", "WELL-2", "WELL-4", "WELL-5"]

    by_key = {(r[0], r[2]): r for r in body}
    # L1: WELL-1 / P-A / 5000 / +30d / Confirmed (the WELL's status) / Primary
    l1 = by_key[("WELL-1", 5000)]
    assert l1[1] == w.p_a_desc
    # openpyxl reads a date-formatted cell back as a datetime at midnight, which is
    # exactly what `_parse_ros` accepts as "a real Excel date cell".
    assert l1[3].date() == (NOW + timedelta(days=30)).date()
    assert l1[4] == "Confirmed"
    assert l1[5] == "Primary"
    # Budgeted and Planned wells export their OWN status, so nothing is normalised
    # away and an unmodified re-upload asserts what is already true.
    assert by_key[("WELL-4", 500)][4] == "Budgeted"
    assert by_key[("WELL-5", 2000)][4] == "Planned"


def test_template_is_scoped_to_the_named_customer_only(client_world):
    """Beta's WELL-3 must not appear in Acme's file. The scope is the customer,
    because that is the scope coverage is computed at."""
    client, _sf, w = client_world
    _t, acme = _sheet_rows(_download(client, w.acme_id).content)
    assert "WELL-3" not in {r[0] for r in acme[1:]}

    _t, beta = _sheet_rows(_download(client, w.beta_id).content)
    assert {r[0] for r in beta[1:]} == {"WELL-3"}


def test_well_id_narrows_further_and_a_foreign_well_is_refused(client_world):
    client, _sf, w = client_world
    _t, rows = _sheet_rows(_download(client, w.acme_id, well_id=w.w2_id).content)
    assert len(rows) == 3  # header + WELL-2's two lines
    assert {r[0] for r in rows[1:]} == {"WELL-2"}

    # Beta's well under Acme's scope: refused, not silently widened or emptied.
    resp = client.get(
        f"/demand-imports/template?customer_id={w.acme_id}&well_id={w.w3_id}"
    )
    assert resp.status_code == 400
    assert "does not belong to customer" in resp.json()["detail"]

    assert client.get("/demand-imports/template?customer_id=nope").status_code == 404
    assert (
        client.get(
            f"/demand-imports/template?customer_id={w.acme_id}&well_id=nope"
        ).status_code
        == 404
    )
    # customer_id is REQUIRED -- there is deliberately no all-customers export.
    assert client.get("/demand-imports/template").status_code == 422


def test_template_carries_its_guidance_on_a_second_sheet_the_parser_ignores(
    client_world,
):
    """The notes live in the workbook, not only on a screen the planner may have
    closed. `_read_rows` reads worksheets[0] only, so the second sheet is inert."""
    client, _sf, w = client_world
    payload = _download(client, w.acme_id).content
    title, notes = _sheet_rows(payload, sheet=1)
    assert title == "Notes"
    prose = " ".join(str(r[0]) for r in notes if r and r[0])
    assert "NOT a blank template" in prose
    assert "5 demand line(s), the complete in-scope book (no row cap)" in prose
    assert "cascades a revision to every line" in prose

    # And it really is inert: uploading the file stages 5 rows, not 5 plus notes.
    assert _upload(client, payload).json()["row_count"] == 5


def test_a_customer_with_no_demand_gets_a_header_only_file_not_a_404(client_world):
    """"This customer's book is empty" is a true answer; 404 ("no such thing")
    would misreport it. Same distinction the rest of the platform draws between an
    absent row and a zero."""
    client, session_factory, w = client_world
    db = session_factory()
    try:
        for line in db.query(DemandLine).all():
            cr = line.coverage_result
            if cr is not None:
                db.delete(cr)
            db.delete(line)
        db.commit()
    finally:
        db.close()

    resp = _download(client, w.acme_id)
    assert resp.headers["X-Demand-Row-Count"] == "0"
    _title, rows = _sheet_rows(resp.content)
    assert rows == [TEMPLATE_COLUMNS]


# --------------------------------------------------------------------------
# THE round trip
# --------------------------------------------------------------------------


def test_generated_template_round_trips_with_no_changes(client_world):
    """Download, re-upload untouched, and NOTHING is reported as changing.

    Every row must come back as an exact-match REVISION whose file values equal the
    matched line's live values on all four fields, with no errors and -- crucially --
    no conflicts: the status cell carries the well's own current status, so it
    asserts nothing new.
    """
    client, session_factory, w = client_world
    revisions_before = _revision_count(session_factory)

    payload = _download(client, w.acme_id).content
    resp = _upload(client, payload, filename="demand-Acme.xlsx")
    assert resp.status_code == 201, resp.text
    batch = resp.json()

    assert batch["row_count"] == 5
    assert batch["error_count"] == 0
    assert batch["revision_suggestion_count"] == 5
    assert batch["new_suggestion_count"] == 0
    assert batch["conflict_count"] == 0
    assert batch["unapproved_conflict_count"] == 0

    for row in batch["rows"]:
        assert row["error"] is None, row
        assert row["match_type"] == "Revision", row
        assert "Exact match" in row["match_reason"], row
        assert row["is_conflict"] is False, row
        assert row["requires_override_approval"] is False, row
        # No change on ANY of the four fields the row can carry.
        assert row["quantity"] == row["matched_quantity"], row
        assert row["ros_date"][:10] == row["matched_ros_date"][:10], row
        assert row["status"] == row["matched_status"], row
        assert row["profile"] == row["matched_profile"], row

    # Staging wrote no demand, as ever.
    assert _revision_count(session_factory) == revisions_before


def test_round_tripping_an_EDITED_template_reports_exactly_the_edited_cell(
    client_world,
):
    """The other half of the guarantee: a template is useful only if changing one
    cell changes one row. Everything else must still stage as an exact match."""
    client, _sf, w = client_world
    payload = _download(client, w.acme_id).content

    wb = load_workbook(io.BytesIO(payload))
    ws = wb.worksheets[0]
    edited_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "WELL-1":
            ws.cell(row=r, column=3).value = 6500
            edited_row = r
            break
    assert edited_row is not None
    buffer = io.BytesIO()
    wb.save(buffer)

    batch = _upload(client, buffer.getvalue()).json()
    assert batch["error_count"] == 0
    assert batch["conflict_count"] == 0
    changed = [
        r for r in batch["rows"] if r["quantity"] != r["matched_quantity"]
    ]
    assert len(changed) == 1
    assert changed[0]["row_number"] == edited_row
    assert changed[0]["matched_demand_line_id"] == w.l1_id
    assert changed[0]["quantity"] == 6500
    assert "quantity differs" in changed[0]["match_reason"]


def test_template_reflects_CURRENT_data_not_a_stale_snapshot(client_world):
    """Generated per request from the live book. A cached or seeded export would
    hand a planner yesterday's numbers to edit, and the edit would silently revert
    somebody else's change."""
    client, session_factory, w = client_world

    _t, before = _sheet_rows(_download(client, w.acme_id).content)
    assert ("WELL-1", w.p_a_desc, 5000) == before[
        [r[0] for r in before].index("WELL-1")
    ][:3]

    # Move the world: a quantity/ROS revision on L1, and a well-status change on
    # WELL-4 -- both of the things the template states.
    db = session_factory()
    try:
        line = db.get(DemandLine, w.l1_id)
        apply_revision(
            db,
            line,
            quantity=7250,
            ros_date=NOW + timedelta(days=61),
            profile=DemandProfile.CONTINGENCY,
        )
        set_well_demand_status(
            db, db.get(Well, w.w4_id), DemandStatus.CONFIRMED
        )
        db.commit()
    finally:
        db.close()

    _t, after = _sheet_rows(_download(client, w.acme_id).content)
    body = {r[0]: r for r in after[1:] if r[0] in ("WELL-1", "WELL-4")}
    assert body["WELL-1"][2] == 7250
    assert body["WELL-1"][3].date() == (NOW + timedelta(days=61)).date()
    assert body["WELL-1"][5] == "Contingency"
    assert body["WELL-4"][4] == "Confirmed"

    # And the FRESH file still round-trips clean against the world it came from.
    batch = _upload(client, _download(client, w.acme_id).content).json()
    assert batch["error_count"] == 0
    assert batch["conflict_count"] == 0
    assert batch["revision_suggestion_count"] == 5


def test_a_profile_only_duplicate_pair_is_WARNED_ABOUT_not_silently_dropped(
    client_world,
):
    """The parser's in-file duplicate check keys on well+product+ROS+quantity and
    NOT profile, so two lines differing only by profile export as an apparent
    duplicate. The export stays faithful and says so in the Notes sheet; dropping a
    line to make the round trip look clean is the one thing it must not do."""
    client, session_factory, w = client_world
    db = session_factory()
    try:
        twin = db.get(DemandLine, w.l1_id)
        db.add(
            DemandLine(
                well_id=twin.well_id,
                product_id=twin.product_id,
                quantity=twin.quantity,
                ros_date=twin.ros_date,
                profile=DemandProfile.CONTINGENCY,
            )
        )
        db.commit()
    finally:
        db.close()

    resp = _download(client, w.acme_id)
    assert resp.headers["X-Demand-Row-Count"] == "6"  # BOTH lines are exported
    _t, notes = _sheet_rows(resp.content, sheet=1)
    prose = " ".join(str(r[0]) for r in notes if r and r[0])
    assert "duplicate check cannot tell apart" in prose
    assert "differ only by profile" in prose

    # And the warning is TRUE -- one row does stage as an error, as predicted.
    batch = _upload(client, resp.content).json()
    assert batch["error_count"] == 1
    dup = next(r for r in batch["rows"] if r["match_type"] == "Error")
    assert "duplicate of row" in dup["error"]


def _revision_count(session_factory) -> int:
    from app.models import DemandRevision

    db = session_factory()
    try:
        return db.query(DemandRevision).count()
    finally:
        db.close()

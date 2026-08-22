"""COMPANY-OWNED inventory maintenance surface.

MVP-COMPROMISE[C-03]: this whole module tests a write path onto tables design
principle #5 calls read-only Oracle projections. See
`app.engines.company_inventory`'s module docstring and MVP_COMPROMISES.md C-03.

Covers, per the task brief:

  * an absent on-hand row reads UNKNOWN, never zero
  * a quantity=0 on-order row reads as "nothing on order", distinct from absent
  * an inline edit stamps source_system="manual" and synced_at
  * a row sourced outside PLATFORM_MAINTAINABLE_SOURCES is refused with 403
  * the on-hand template contains the CURRENT position
  * an upload replaces the named position; bad rows are reported, not dropped
  * an on-hand edit that drops a line below demand flips a coverage verdict, and
    the response reports the flip
  * a write in one Business Unit never touches another BU's inventory or verdicts
  * a cross-product aggregate (assignments across two products of different units)
    is reported via quantities_by_unit with a null scalar
"""

import io
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.db import Base, get_db
from app.engines.company_inventory import (
    NotMaintainable,
    PLATFORM_MAINTAINABLE_SOURCES,
    build_on_hand_template,
    create_on_hand,
    get_position,
    parse_and_replace_on_hand,
    set_on_hand,
)
from app.engines.executive import quantity_by_unit
from app.main import app
from app.models import (
    AllocationPolicy,
    BusinessUnit,
    Customer,
    DemandLine,
    DemandProfile,
    DemandStatus,
    InventoryAssignment,
    InventoryOnHand,
    InventoryOnOrder,
    PlanningNode,
    Product,
    UnitOfMeasure,
    Well,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _bu(db, name="Company BU"):
    bu = BusinessUnit(name=name)
    db.add(bu)
    db.flush()
    return bu


def _customer(db, bu, name, policy=AllocationPolicy.SOFT):
    customer = Customer(name=name, allocation_policy=policy, business_unit_id=bu.id)
    db.add(customer)
    db.flush()
    node = PlanningNode(
        customer_id=customer.id, node_type="Campaign", name=f"{name} campaign"
    )
    db.add(node)
    db.flush()
    return customer, node


def _product(db, description, unit=UnitOfMeasure.MTR):
    product = Product(
        type="CSG", size="9-5/8", weight=53.5, grade="P110", grade_type="Carbon",
        connection="VAM 21", commodity="SMLS",
        description=description, unit_of_measure=unit,
    )
    db.add(product)
    db.flush()
    return product


def _on_hand(db, bu, product, quantity, source_system="synthetic"):
    row = InventoryOnHand(
        business_unit_id=bu.id, product_id=product.id, quantity=quantity,
        source_system=source_system,
    )
    db.add(row)
    db.flush()
    return row


def _well(db, node, name, status=DemandStatus.CONFIRMED):
    well = Well(planning_node_id=node.id, name=name, demand_status=status)
    db.add(well)
    db.flush()
    return well


def _line(db, well, product, quantity, days_out=30):
    line = DemandLine(
        well_id=well.id, product_id=product.id, quantity=quantity,
        ros_date=datetime.utcnow() + timedelta(days=days_out),
        profile=DemandProfile.PRIMARY,
    )
    db.add(line)
    db.flush()
    return line


def _sheet(rows, headers=("product", "on_hand_quantity")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "On-Hand"
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def client(db_session):
    def _override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


# ---------------------------------------------------------------------------
# Unknown vs. zero
# ---------------------------------------------------------------------------


def test_absent_on_hand_row_reads_unknown_not_zero(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "No on-hand row product")
    # Nothing written -- but reference the product via an on-order row so it is
    # in scope for the position read (see get_position's docstring).
    db_session.add(
        InventoryOnOrder(business_unit_id=bu.id, product_id=product.id, quantity=5)
    )
    db_session.flush()

    position = get_position(db_session, bu.id)
    on_hand = next(r for r in position.on_hand if r.product_id == product.id)
    assert on_hand.known is False
    assert on_hand.quantity is None


def test_explicit_zero_on_order_row_is_nothing_on_order_not_absent(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Zero on-order product")
    db_session.add(
        InventoryOnOrder(business_unit_id=bu.id, product_id=product.id, quantity=0)
    )
    db_session.flush()

    position = get_position(db_session, bu.id)
    rows = [r for r in position.on_order if r.product_id == product.id]
    assert len(rows) == 1
    assert rows[0].quantity == 0.0  # measured zero, a row exists


def test_product_with_no_data_anywhere_is_absent_from_the_unscoped_list(db_session):
    bu = _bu(db_session)
    _product(db_session, "Untouched product")  # no rows in this BU at all
    position = get_position(db_session, bu.id)
    assert position.on_hand == ()
    assert position.on_order == ()


# ---------------------------------------------------------------------------
# The source_system gate
# ---------------------------------------------------------------------------


def test_inline_edit_stamps_manual_and_synced_at(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Editable product")
    row = _on_hand(db_session, bu, product, 100, source_system="synthetic")

    before = datetime.utcnow()
    result = set_on_hand(db_session, row.id, 250)
    db_session.flush()

    assert result.after["quantity"] == 250
    assert result.after["source_system"] == "manual"
    assert result.after["synced_at"] >= before
    refreshed = db_session.get(InventoryOnHand, row.id)
    assert refreshed.quantity == 250
    assert refreshed.source_system == "manual"
    assert refreshed.synced_at is not None


def test_row_sourced_outside_maintainable_set_is_refused(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Oracle-fed product")
    row = _on_hand(db_session, bu, product, 100, source_system="oracle")
    assert "oracle" not in PLATFORM_MAINTAINABLE_SOURCES

    with pytest.raises(NotMaintainable) as excinfo:
        set_on_hand(db_session, row.id, 999)

    message = str(excinfo.value)
    assert row.id in message
    assert "oracle" in message
    # Unaffected: the refusal must not have written anything (raised before any
    # mutation, so there is nothing to roll back).
    refreshed = db_session.get(InventoryOnHand, row.id)
    assert refreshed.quantity == 100
    assert refreshed.source_system == "oracle"


def test_row_sourced_outside_maintainable_set_is_403_at_the_api(client, db_session):
    bu = _bu(db_session)
    product = _product(db_session, "API oracle product")
    row = _on_hand(db_session, bu, product, 100, source_system="oracle")

    resp = client.patch(
        f"/company-inventory/on-hand/{row.id}", json={"quantity": 999}
    )
    assert resp.status_code == 403
    assert "oracle" in resp.json()["detail"]


# ---------------------------------------------------------------------------
# Template download
# ---------------------------------------------------------------------------


def test_template_contains_the_current_position(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Templated product")
    _on_hand(db_session, bu, product, 777, source_system="synthetic")

    template = build_on_hand_template(db_session, bu)
    assert template.row_count == 1

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(template.content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = [tuple(r) for r in sheet.iter_rows(values_only=True)]
    workbook.close()
    assert rows[0] == ("product", "on_hand_quantity")
    assert rows[1] == ("Templated product", 777)


# ---------------------------------------------------------------------------
# Upload replaces the named position; bad rows are reported, not dropped
# ---------------------------------------------------------------------------


def test_upload_replaces_named_products_and_reports_bad_rows(db_session):
    bu = _bu(db_session)
    known = _product(db_session, "Known upload product")
    untouched = _product(db_session, "Untouched by upload")
    _on_hand(db_session, bu, known, 10, source_system="synthetic")
    _on_hand(db_session, bu, untouched, 55, source_system="synthetic")

    payload = _sheet(
        [
            ("Known upload product", 500),
            ("Nonexistent product", 20),
            ("Known upload product", -5),  # duplicate AND negative -- reported once
        ]
    )
    result = parse_and_replace_on_hand(db_session, bu, payload, filename="t.xlsx")

    assert result.replaced_count == 1
    assert result.error_count == 2
    errors = [r for r in result.rows if r.action == "Error"]
    assert any("unknown product" in (e.error or "") for e in errors)

    known_row = (
        db_session.query(InventoryOnHand)
        .filter(InventoryOnHand.business_unit_id == bu.id, InventoryOnHand.product_id == known.id)
        .one()
    )
    assert known_row.quantity == 500
    assert known_row.source_system == "manual"

    # Untouched product's row is unchanged -- replace only what the file names.
    untouched_row = (
        db_session.query(InventoryOnHand)
        .filter(InventoryOnHand.business_unit_id == bu.id, InventoryOnHand.product_id == untouched.id)
        .one()
    )
    assert untouched_row.quantity == 55
    assert untouched_row.source_system == "synthetic"


def test_upload_row_targeting_an_oracle_sourced_row_is_a_row_level_error(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Oracle row for upload")
    _on_hand(db_session, bu, product, 42, source_system="oracle")

    payload = _sheet([("Oracle row for upload", 999)])
    result = parse_and_replace_on_hand(db_session, bu, payload, filename="t.xlsx")

    assert result.error_count == 1
    assert result.replaced_count == 0
    row = db_session.get(InventoryOnHand, (
        db_session.query(InventoryOnHand)
        .filter(InventoryOnHand.business_unit_id == bu.id)
        .one()
    ).id)
    assert row.quantity == 42  # untouched


# ---------------------------------------------------------------------------
# Coverage recomputation on write
# ---------------------------------------------------------------------------


def test_on_hand_edit_below_demand_flips_coverage_and_is_reported(db_session):
    bu = _bu(db_session)
    customer, node = _customer(db_session, bu, "Coverage Co")
    product = _product(db_session, "Coverage product")
    row = _on_hand(db_session, bu, product, 1000, source_system="synthetic")
    well = _well(db_session, node, "Coverage well")
    _line(db_session, well, product, 800)

    from app.engines.coverage import recompute_customer

    recompute_customer(db_session, customer)
    db_session.flush()
    assert well.coverage_status == "Covered"

    result = set_on_hand(db_session, row.id, 100)  # now short of the 800 line
    db_session.flush()

    assert well.coverage_status == "Uncovered"
    assert well.id in result.recompute.well_changes
    before, after = result.recompute.well_changes[well.id]
    assert before == "Covered"
    assert after == "Uncovered"


# ---------------------------------------------------------------------------
# Business Unit boundary
# ---------------------------------------------------------------------------


def test_write_in_one_bu_never_touches_another_bus_inventory_or_verdicts(db_session):
    bu_a = _bu(db_session, "BU A")
    bu_b = _bu(db_session, "BU B")
    customer_a, node_a = _customer(db_session, bu_a, "Customer A")
    customer_b, node_b = _customer(db_session, bu_b, "Customer B")
    product = _product(db_session, "Shared catalogue product")

    row_a = _on_hand(db_session, bu_a, product, 1000, source_system="synthetic")
    row_b = _on_hand(db_session, bu_b, product, 1000, source_system="synthetic")
    well_a = _well(db_session, node_a, "Well A")
    well_b = _well(db_session, node_b, "Well B")
    _line(db_session, well_a, product, 800)
    _line(db_session, well_b, product, 800)

    from app.engines.coverage import recompute_customer

    recompute_customer(db_session, customer_a)
    recompute_customer(db_session, customer_b)
    db_session.flush()
    assert well_a.coverage_status == "Covered"
    assert well_b.coverage_status == "Covered"

    result = set_on_hand(db_session, row_a.id, 50)
    db_session.flush()

    assert well_a.coverage_status == "Uncovered"
    assert well_b.coverage_status == "Covered"  # untouched
    assert well_b.id not in result.recompute.well_changes
    # BU B's own on-hand row is untouched.
    refreshed_b = db_session.get(InventoryOnHand, row_b.id)
    assert refreshed_b.quantity == 1000
    assert refreshed_b.source_system == "synthetic"


# ---------------------------------------------------------------------------
# quantity_by_unit convention for cross-product aggregation
# ---------------------------------------------------------------------------


def test_quantities_aggregated_across_units_report_breakdown_and_null_scalar(db_session):
    bu = _bu(db_session)
    customer, node = _customer(db_session, bu, "Units Co")
    product_mtr = _product(db_session, "Metres product", unit=UnitOfMeasure.MTR)
    product_mt = _product(db_session, "Tonnes product", unit=UnitOfMeasure.MT)
    well = _well(db_session, node, "Coverage well")
    line_mtr = _line(db_session, well, product_mtr, 100)
    line_mt = _line(db_session, well, product_mt, 50)

    db_session.add(
        InventoryAssignment(
            demand_line_id=line_mtr.id, product_id=product_mtr.id, quantity=40,
            source_system="synthetic",
        )
    )
    db_session.add(
        InventoryAssignment(
            demand_line_id=line_mt.id, product_id=product_mt.id, quantity=20,
            source_system="synthetic",
        )
    )
    db_session.flush()

    position = get_position(db_session, bu.id)
    pairs = [
        (group.unit_of_measure, group.total_quantity) for group in position.assignments
    ]
    breakdown, single = quantity_by_unit(pairs)
    assert single is None  # two units contributed -- no safe scalar
    units = {b.unit_of_measure for b in breakdown}
    assert units == {UnitOfMeasure.MTR, UnitOfMeasure.MT}


# ---------------------------------------------------------------------------
# Create / delete
# ---------------------------------------------------------------------------


def test_create_on_hand_where_none_exists(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Brand new row product")
    result = create_on_hand(db_session, bu.id, product.id, 321)
    db_session.flush()

    row = (
        db_session.query(InventoryOnHand)
        .filter(InventoryOnHand.business_unit_id == bu.id, InventoryOnHand.product_id == product.id)
        .one()
    )
    assert row.quantity == 321
    assert row.source_system == "manual"
    assert result.before == {}


def test_create_on_hand_refuses_when_a_row_already_exists(db_session):
    bu = _bu(db_session)
    product = _product(db_session, "Already has a row")
    _on_hand(db_session, bu, product, 5, source_system="synthetic")

    with pytest.raises(ValueError):
        create_on_hand(db_session, bu.id, product.id, 999)

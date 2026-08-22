"""Template download: the CURRENT declared customer-owned position, re-uploadable.

THE CORRECTNESS BAR, and it is a round trip
-------------------------------------------
`GET /customer-owned-inventory/{customer_id}/template` is not a blank column hint. It
is an export of the live declared position, and its whole value rests on one property:

    re-uploading it UNMODIFIED must restate every position at the same quantity --
    every row reported as "Replaced" with previous_quantity == quantity, no errors,
    and no coverage verdict moving.

"Replaced with the same number" is this flow's spelling of "nothing changed": there is
no Created-vs-Replaced question to get wrong, because every exported row is a position
that already exists (see `app.engines.customer_owned_import` on why there is no
staging step). A planner edits the cells they mean to change; nothing else may move.

That property is a statement about the GENERATOR and the PARSER agreeing, so it is not
checkable by inspecting either one. It is tested by actually generating a file,
actually posting it back through the real upload endpoint, and asserting the reported
result says nothing changed. `test_generated_template_round_trips_with_no_changes` is
the load-bearing test in this module.

AND THE TWO WAYS OF BEING EMPTY ARE KEPT APART
---------------------------------------------
A never-uploaded customer and a customer who uploaded and declares nothing both get a
header-only workbook -- deliberately the SAME file shape, because the workbook is a
form and a form must not change shape depending on which state you are in. They are
distinguished by the `X-Customer-Owned-Has-Uploaded` header and by the Notes sheet's
prose, not by the columns. Row count alone cannot tell them apart, which is exactly
why the flag is a separate header.
"""

import io
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import get_db
from app.engines.coverage import recompute_customer
from app.engines.customer_owned_import import TEMPLATE_COLUMNS
from app.main import app
from app.models import CustomerOwnedInventory, UnitOfMeasure
from tests.test_customer_owned_inventory import (
    _bu,
    _customer,
    _line,
    _owns,
    _product,
    _sheet,
    _stock,
    _upload_row,
    _well,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def client(db_session):
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


def _download(client, customer_id):
    resp = client.get(f"/customer-owned-inventory/{customer_id}/template")
    assert resp.status_code == 200, resp.text
    return resp


def _sheet_rows(payload: bytes, sheet=0):
    wb = load_workbook(io.BytesIO(payload), data_only=True)
    ws = wb.worksheets[sheet]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    title = ws.title
    wb.close()
    return title, rows


def _reupload(client, customer_id, payload, filename="edited.xlsx"):
    return client.post(
        f"/customer-owned-inventory/{customer_id}/uploads",
        files={"file": (filename, payload, XLSX)},
    )


def _world(db):
    """One customer with TWO declared positions, plus a well so coverage is real."""
    bu = _bu(db)
    customer, node = _customer(db, bu, "Template Co")
    p1 = _product(db, "CSG 9-5/8 template alpha")
    p2 = _product(db, "CSG 9-5/8 template beta")
    _stock(db, bu, p1, 0)
    _stock(db, bu, p2, 0)
    well = _well(db, node, "TPL-WELL-1")
    _line(db, well, p1, 1000, 30)
    upload = _upload_row(db, customer, filename="first-count.xlsx")
    _owns(db, customer, p1, 4200, upload=upload)
    _owns(db, customer, p2, 750.5, upload=upload)
    # Provenance as a real upload would leave it -- the composite `_reference` string
    # that the template deliberately does NOT put back in the `location` column.
    for row in db.query(CustomerOwnedInventory).filter(
        CustomerOwnedInventory.customer_id == customer.id
    ):
        row.source_reference = "first-count.xlsx"
    # Coverage computed ONCE up front, so the stored rollups exist. Without this the
    # very first upload of the test would report every well as moving from None, which
    # is the fixture being fresh rather than the template changing anything.
    recompute_customer(db, customer)
    db.commit()
    return customer, p1, p2, well


# --------------------------------------------------------------------------
# Shape
# --------------------------------------------------------------------------


def test_template_is_a_real_xlsx_with_the_canonical_header(client, db_session):
    customer, _p1, _p2, _w = _world(db_session)
    resp = _download(client, customer.id)

    assert resp.headers["content-type"].startswith(XLSX)
    assert "attachment;" in resp.headers["content-disposition"]
    assert "customer-owned-Template-Co.xlsx" in resp.headers["content-disposition"]
    # So a client can caption "2 declared positions" without opening the workbook.
    assert resp.headers["X-Customer-Owned-Row-Count"] == "2"
    assert resp.headers["X-Customer-Owned-Has-Uploaded"] == "true"

    title, rows = _sheet_rows(resp.content)
    assert title == "Customer-Owned Inventory"
    # The header is the CANONICAL contract, in order -- not an alias spelling that a
    # later edit to `_ALIASES` could silently break.
    assert rows[0] == TEMPLATE_COLUMNS
    assert list(TEMPLATE_COLUMNS) == [
        "product",
        "quantity",
        "customer",
        "as_of_date",
        "location",
    ]


def test_template_body_is_the_customers_current_position_one_row_per_product(
    client, db_session
):
    """Two declared positions, both exported, product written as its DESCRIPTION."""
    customer, p1, p2, _w = _world(db_session)
    _title, rows = _sheet_rows(_download(client, customer.id).content)
    body = rows[1:]
    assert len(body) == 2

    by_product = {r[0]: r for r in body}
    assert set(by_product) == {p1.description, p2.description}

    alpha = by_product[p1.description]
    assert alpha[1] == 4200
    assert by_product[p2.description][1] == 750.5
    # The `customer` cell is filled with the NAME, so uploading against the wrong
    # operator is refused row by row rather than silently succeeding.
    assert alpha[2] == "Template Co"
    # `as_of_date` is a real date cell -- openpyxl reads it back as midnight, which is
    # exactly what `_parse_as_of` accepts as "a real Excel date cell".
    assert alpha[3].date() == datetime.utcnow().date()
    # `location` is deliberately blank: there is no location field on a position, and
    # the composite `source_reference` cannot be split back apart.
    assert alpha[4] is None


def test_template_is_scoped_to_the_named_customer_only(client, db_session):
    """Another operator's declared stock must never appear in this file -- the
    tightest boundary in the platform: this material is that customer's PROPERTY."""
    customer, p1, _p2, _w = _world(db_session)
    bu = _bu(db_session)
    other, _node = _customer(db_session, bu, "Neighbour Co")
    other_product = _product(db_session, "CSG neighbour only")
    _owns(db_session, other, other_product, 9999)
    db_session.commit()

    _t, mine = _sheet_rows(_download(client, customer.id).content)
    assert other_product.description not in {r[0] for r in mine[1:]}

    _t, theirs = _sheet_rows(_download(client, other.id).content)
    assert {r[0] for r in theirs[1:]} == {other_product.description}
    assert p1.description not in {r[0] for r in theirs[1:]}


def test_an_unknown_customer_is_a_404(client, db_session):
    assert (
        client.get("/customer-owned-inventory/no-such-customer/template").status_code
        == 404
    )


def test_template_carries_its_guidance_on_a_second_sheet_the_parser_ignores(
    client, db_session
):
    """The notes live in the workbook, not only on a screen the planner may have
    closed. `_read_rows` reads worksheets[0] only, so the second sheet is inert."""
    customer, _p1, _p2, _w = _world(db_session)
    payload = _download(client, customer.id).content

    title, notes = _sheet_rows(payload, sheet=1)
    assert title == "Notes"
    prose = " ".join(str(r[0]) for r in notes if r and r[0])
    assert "NOT a blank template" in prose
    assert "2 declared position(s), the complete position (no row cap)" in prose
    # The two traps a planner can walk into, both stated.
    assert "Deleting a row therefore does NOT set that position to zero" in prose
    assert "location` is deliberately left BLANK" in prose
    # The provenance that is NOT put in the location column is still reproduced.
    assert "first-count.xlsx" in prose

    # And it really is inert: uploading the file reports 2 rows, not 2 plus notes.
    assert _reupload(client, customer.id, payload).json()["row_count"] == 2


# --------------------------------------------------------------------------
# THE round trip
# --------------------------------------------------------------------------


def test_generated_template_round_trips_with_no_changes(client, db_session):
    """Download, re-upload untouched, and NOTHING is reported as changing.

    Every row comes back as a "Replaced" whose previous quantity equals its new one,
    with no errors, no positions created, and no well's coverage verdict moving.
    """
    customer, p1, p2, _w = _world(db_session)
    payload = _download(client, customer.id).content

    resp = _reupload(
        client, customer.id, payload, filename="customer-owned-Template-Co.xlsx"
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()

    assert body["row_count"] == 2
    assert body["error_count"] == 0
    assert body["created_count"] == 0
    assert body["replaced_count"] == 2
    assert body["applied_count"] == 2
    # Not one well moved: the quantities re-asserted are the ones already stored.
    assert body["coverage_changes"] == {}

    for row in body["rows"]:
        assert row["error"] is None, row
        assert row["action"] == "Replaced", row
        # THE bar: the position was restated at exactly the value it already held.
        assert row["previous_quantity"] == row["quantity"], row
        assert row["unit_of_measure"] == UnitOfMeasure.MTR.value, row

    quantities = {r["product_description"]: r["quantity"] for r in body["rows"]}
    assert quantities == {p1.description: 4200, p2.description: 750.5}

    # And the stored position is unchanged in the database, not merely reported so.
    stored = {
        row.product_id: row.quantity
        for row in db_session.query(CustomerOwnedInventory).filter(
            CustomerOwnedInventory.customer_id == customer.id
        )
    }
    assert stored == {p1.id: 4200, p2.id: 750.5}


def test_the_round_trip_preserves_the_AS_OF_DATE_rather_than_stamping_today(
    client, db_session
):
    """A stale count must not be made to look current by a round trip.

    The `as_of_date` column is written precisely so this cannot happen. The one honest
    cost, asserted rather than hidden: the stored time of day is truncated to midnight
    of the SAME date, because the cell is a date. The date is what staleness is judged
    on, and it survives exactly.
    """
    bu = _bu(db_session)
    customer, _node = _customer(db_session, bu, "Stale Count Co")
    product = _product(db_session, "CSG stale count")
    _stock(db_session, bu, product, 0)
    upload = _upload_row(db_session, customer)
    old = datetime.utcnow() - timedelta(days=95)
    db_session.add(
        CustomerOwnedInventory(
            customer_id=customer.id,
            product_id=product.id,
            quantity=1234,
            source_system="customer-upload",
            uploaded_at=old,
            upload_id=upload.id,
        )
    )
    db_session.commit()

    payload = _download(client, customer.id).content
    _t, rows = _sheet_rows(payload)
    assert rows[1][3].date() == old.date()

    assert _reupload(client, customer.id, payload).json()["error_count"] == 0

    row = (
        db_session.query(CustomerOwnedInventory)
        .filter(CustomerOwnedInventory.customer_id == customer.id)
        .one()
    )
    assert row.quantity == 1234
    # The DATE is preserved -- a 95-day-old count still reads as 95 days old.
    assert row.uploaded_at.date() == old.date()
    # Midnight, because the column is a date. Documented in the Notes sheet.
    assert (row.uploaded_at.hour, row.uploaded_at.minute) == (0, 0)


def test_round_tripping_an_EDITED_template_changes_exactly_the_edited_cell(
    client, db_session
):
    """The other half of the guarantee: a template is useful only if changing one
    cell changes one position. Everything else must restate unchanged."""
    customer, p1, p2, _w = _world(db_session)
    payload = _download(client, customer.id).content

    wb = load_workbook(io.BytesIO(payload))
    ws = wb.worksheets[0]
    edited_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == p1.description:
            ws.cell(row=r, column=2).value = 6500
            edited_row = r
            break
    assert edited_row is not None
    buffer = io.BytesIO()
    wb.save(buffer)

    body = _reupload(client, customer.id, buffer.getvalue()).json()
    assert body["error_count"] == 0
    assert body["created_count"] == 0
    changed = [r for r in body["rows"] if r["previous_quantity"] != r["quantity"]]
    assert len(changed) == 1
    assert changed[0]["row_number"] == edited_row
    assert changed[0]["product_id"] == p1.id
    assert changed[0]["quantity"] == 6500
    assert changed[0]["previous_quantity"] == 4200

    untouched = next(r for r in body["rows"] if r["product_id"] == p2.id)
    assert untouched["quantity"] == untouched["previous_quantity"] == 750.5


def test_template_reflects_the_CURRENT_position_not_a_stale_snapshot(
    client, db_session
):
    """Generated per request from the live position. A cached export would hand a
    planner yesterday's numbers to edit, and the edit would silently revert somebody
    else's upload."""
    customer, p1, p2, _w = _world(db_session)

    _t, before = _sheet_rows(_download(client, customer.id).content)
    assert {r[0]: r[1] for r in before[1:]} == {
        p1.description: 4200,
        p2.description: 750.5,
    }

    # Move the world THROUGH THE API -- a new upload restating one product and adding
    # a third. This is exactly what a colleague doing their own edit looks like.
    p3 = _product(db_session, "CSG 9-5/8 template gamma")
    _stock(db_session, _bu(db_session), p3, 0)
    db_session.commit()
    assert (
        _reupload(
            client,
            customer.id,
            _sheet([(p1.description, 8800), (p3.description, 15)]),
            filename="colleague.xlsx",
        ).status_code
        == 201
    )

    resp = _download(client, customer.id)
    assert resp.headers["X-Customer-Owned-Row-Count"] == "3"
    _t, after = _sheet_rows(resp.content)
    assert {r[0]: r[1] for r in after[1:]} == {
        p1.description: 8800,   # restated by the colleague's upload
        p2.description: 750.5,  # untouched -- an upload replaces only what it names
        p3.description: 15,     # created by the colleague's upload
    }

    # And the FRESH file still round-trips clean against the world it came from.
    body = _reupload(client, customer.id, resp.content).json()
    assert body["error_count"] == 0
    assert body["created_count"] == 0
    assert body["replaced_count"] == 3
    assert all(r["previous_quantity"] == r["quantity"] for r in body["rows"])


# --------------------------------------------------------------------------
# The two ways of being empty -- same file shape, different fact
# --------------------------------------------------------------------------


def test_a_NEVER_UPLOADED_customer_gets_a_valid_header_only_template(
    client, db_session
):
    """"Nobody has ever told us" is a true answer about a customer that exists; a
    404 ("no such thing") would misreport it. And this is the state the download
    matters MOST in -- there is nothing on screen to review, so the file is the only
    way to start."""
    bu = _bu(db_session)
    customer, _node = _customer(db_session, bu, "Never Uploaded Co")
    db_session.commit()

    # Precondition, from the platform's own discriminator.
    position = client.get(f"/customer-owned-inventory/{customer.id}").json()
    assert position["has_uploaded"] is False

    resp = _download(client, customer.id)
    assert resp.headers["X-Customer-Owned-Row-Count"] == "0"
    assert resp.headers["X-Customer-Owned-Has-Uploaded"] == "false"

    # A real, openable workbook with the full header -- not an error, not a stub.
    title, rows = _sheet_rows(resp.content)
    assert title == "Customer-Owned Inventory"
    assert rows == [TEMPLATE_COLUMNS]

    _t, notes = _sheet_rows(resp.content, sheet=1)
    prose = " ".join(str(r[0]) for r in notes if r and r[0])
    assert "NO CUSTOMER-OWNED INVENTORY HAS EVER BEEN UPLOADED" in prose
    assert "NOT the same as the customer owning none" in prose
    # The one consequence a planner can walk into, warned about by name.
    assert "uploading this file UNCHANGED is not a no-op" in prose

    # And it is genuinely re-uploadable: the empty file parses, it does not 400.
    resp = _reupload(client, customer.id, resp.content)
    assert resp.status_code == 201, resp.text
    assert resp.json()["row_count"] == 0
    assert resp.json()["error_count"] == 0


def test_an_UPLOADED_BUT_EMPTY_customer_gets_the_same_shape_and_a_different_note(
    client, db_session
):
    """Uploaded and declaring nothing is a MEASURED FACT, not missing data.

    Same file shape as the never-uploaded case on purpose -- the workbook is a form,
    and a form must not change columns depending on which state you are in. The
    difference is a fact about the customer, so it lives in the header flag and the
    Notes prose. Row count alone cannot tell the two apart, which is precisely why
    `X-Customer-Owned-Has-Uploaded` exists as a separate header rather than being
    inferred from "0 rows".
    """
    bu = _bu(db_session)
    customer, _node = _customer(db_session, bu, "Owns Nothing Co")
    _upload_row(db_session, customer, filename="empty-count.xlsx")
    db_session.commit()

    position = client.get(f"/customer-owned-inventory/{customer.id}").json()
    assert position["has_uploaded"] is True
    assert position["positions"] == []

    resp = _download(client, customer.id)
    # IDENTICAL row count to the never-uploaded case -- so the count is not the
    # discriminator, and the flag is.
    assert resp.headers["X-Customer-Owned-Row-Count"] == "0"
    assert resp.headers["X-Customer-Owned-Has-Uploaded"] == "true"

    title, rows = _sheet_rows(resp.content)
    assert title == "Customer-Owned Inventory"
    assert rows == [TEMPLATE_COLUMNS]

    _t, notes = _sheet_rows(resp.content, sheet=1)
    prose = " ".join(str(r[0]) for r in notes if r and r[0])
    assert "HAS UPLOADED A POSITION AND IT DECLARES NO OWNED MATERIAL AT ALL" in prose
    assert "measured fact, not missing data" in prose
    assert "a different state from 'never uploaded'" in prose
    # And it does NOT claim the other thing.
    assert "HAS EVER BEEN UPLOADED" not in prose

    assert _reupload(client, customer.id, resp.content).status_code == 201

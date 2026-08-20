import io

import openpyxl

from app.models import Customer, CustomerOwnedInventory, CustomerOwnedUpload, Product, UnitOfMeasure


def _fixtures(db):
    c1 = Customer(name="C1")
    c2 = Customer(name="C2")
    p = Product(name="P1", unit_of_measure=UnitOfMeasure.PC)
    db.add_all([c1, c2, p])
    db.commit()
    return c1, c2, p


def _xlsx_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Product", "Quantity", "Unit"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def test_get_unuploaded_customer_returns_has_uploaded_false(client, db):
    c1, c2, p = _fixtures(db)
    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    assert resp.status_code == 200
    body = resp.json()
    assert body["has_uploaded"] is False
    assert body["uploaded_at"] is None
    assert body["positions"] == []


def test_upload_transitions_has_uploaded_to_true(client, db):
    c1, c2, p = _fixtures(db)
    data = _xlsx_bytes([[p.name, 10, "PC"]])

    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200

    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    body = resp.json()
    assert body["has_uploaded"] is True
    assert body["uploaded_at"] is not None
    assert len(body["positions"]) == 1
    assert body["positions"][0]["unit"] == "PC"
    assert body["positions"][0]["quantity"] == 10


def test_upload_empty_file_is_uploaded_but_zero_positions(client, db):
    c1, c2, p = _fixtures(db)
    data = _xlsx_bytes([])

    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200

    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    body = resp.json()
    assert body["has_uploaded"] is True
    assert body["positions"] == []


def test_second_upload_replaces_positions(client, db):
    c1, c2, p = _fixtures(db)
    db.add(CustomerOwnedInventory(customer_id=c1.id, product_id=p.id, quantity=99, unit=UnitOfMeasure.PC))
    db.add(CustomerOwnedUpload(customer_id=c1.id, filename="old.xlsx", row_count=1))
    db.commit()

    data = _xlsx_bytes([[p.name, 5, "PC"]])
    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200

    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    positions = resp.json()["positions"]
    assert len(positions) == 1
    assert positions[0]["quantity"] == 5


def test_upload_validation_errors_include_row_numbers(client, db):
    c1, c2, p = _fixtures(db)
    data = _xlsx_bytes(
        [
            ["Unknown Product", 5, "PC"],
            [p.name, -1, "PC"],
            [p.name, 5, "BadUnit"],
        ]
    )

    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert isinstance(detail, list)
    assert len(detail) == 3
    # rows 2, 3, 4 (row 1 is header)
    row_numbers = {d["row"] for d in detail}
    assert row_numbers == {2, 3, 4}

    # nothing was written since validation failed
    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    assert resp.json()["has_uploaded"] is False


def test_upload_duplicate_product_returns_422_naming_both_rows(client, db):
    c1, c2, p = _fixtures(db)
    data = _xlsx_bytes(
        [
            [p.name, 5, "PC"],
            [p.name, 7, "PC"],
        ]
    )

    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert isinstance(detail, list)
    # row 3 (second occurrence) is flagged as a duplicate, naming row 2 (first occurrence)
    row3 = next(d for d in detail if d["row"] == 3)
    assert any("duplicate" in m.lower() and "2" in m for m in row3["errors"])

    # nothing was written since validation failed
    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    body = resp.json()
    assert body["has_uploaded"] is False
    assert body["positions"] == []


def test_upload_non_numeric_quantity_returns_422_naming_row(client, db):
    c1, c2, p = _fixtures(db)
    data = _xlsx_bytes(
        [
            [p.name, "abc", "PC"],
        ]
    )

    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert isinstance(detail, list)
    assert len(detail) == 1
    assert detail[0]["row"] == 2
    assert any("quantity" in m.lower() for m in detail[0]["errors"])

    # nothing was written since validation failed
    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    body = resp.json()
    assert body["has_uploaded"] is False
    assert body["positions"] == []


def test_upload_isolated_between_customers(client, db):
    c1, c2, p = _fixtures(db)
    data = _xlsx_bytes([[p.name, 10, "PC"]])
    client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("positions.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    resp = client.get(f"/customer-owned-inventory?customer_id={c2.id}")
    body = resp.json()
    assert body["has_uploaded"] is False
    assert body["positions"] == []


def test_template_downloads_and_round_trips_through_upload(client, db):
    c1, c2, p = _fixtures(db)

    resp = client.get("/customer-owned-inventory/template")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == ["Product", "Quantity", "Unit"]

    ws.append([p.name, 12, "PC"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = client.post(
        f"/customer-owned-inventory/upload?customer_id={c1.id}",
        files={"file": ("filled.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200

    resp = client.get(f"/customer-owned-inventory?customer_id={c1.id}")
    positions = resp.json()["positions"]
    assert len(positions) == 1
    assert positions[0]["quantity"] == 12
    assert positions[0]["unit"] == "PC"
